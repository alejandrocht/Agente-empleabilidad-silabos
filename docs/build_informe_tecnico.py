from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\I13305\Desktop\Agente-empleabilidad-silabos")
REFERENCE = Path(r"C:\Users\I13305\Downloads\Segundo Informe.docx")
OUTPUT = ROOT / "docs" / "Informe_tecnico_brechas_CIAR.docx"
FIG_DIR = ROOT / "docs"
CSV_SOURCE = Path(r"C:\Users\I13305\Downloads\neo4j_query_table_data_2026-10-5.csv")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9D9D9", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Calibri", size=10.5, bold=False, italic=False, color="000000") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, space_after=6, line_spacing=1.08) -> None:
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line_spacing


def add_body(doc, text: str, *, bold_lead: str | None = None, space_after=6) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    style_paragraph(paragraph, space_after=space_after)
    if bold_lead and text.startswith(bold_lead):
        set_run_font(paragraph.add_run(bold_lead), bold=True)
        set_run_font(paragraph.add_run(text[len(bold_lead) :]))
    else:
        set_run_font(paragraph.add_run(text))


def add_bullet(doc, text: str, level=0) -> None:
    paragraph = doc.add_paragraph(style="List Paragraph")
    paragraph.paragraph_format.left_indent = Inches(0.22 + level * 0.22)
    paragraph.paragraph_format.first_line_indent = Inches(-0.12)
    style_paragraph(paragraph, space_after=3, line_spacing=1.04)
    set_run_font(paragraph.add_run("• " + text))


def add_heading(doc, text: str, level=1) -> None:
    style_name = "Heading 1" if level == 1 else "Heading 2"
    paragraph = doc.add_paragraph(style=style_name)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_run_font(run, size=16 if level == 1 else 12.5, bold=True)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header = table.rows[0]
    repeat_header(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        cell.text = ""
        set_cell_shading(cell, "000000")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run = cell.paragraphs[0].add_run(value)
        set_run_font(run, size=9, bold=True, color="FFFFFF")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if widths:
            cell.width = Inches(widths[i])
    for row_index, row_values in enumerate(rows):
        row = table.add_row()
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            cell.text = ""
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, "F2F2F2")
            run = cell.paragraphs[0].add_run(value)
            set_run_font(run, size=8.8)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_caption(doc, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    set_run_font(run, size=8.5, italic=True, color="555555")


def add_figure(doc, path: Path, caption: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(path), width=Inches(5.55))
    inline_shape._inline.docPr.set("descr", caption)
    inline_shape._inline.docPr.set("title", caption)
    add_caption(doc, caption)


def add_query(doc, query: str, label: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.0
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F3F3")
    p_pr.append(shd)
    label_run = paragraph.add_run(label + "\n")
    set_run_font(label_run, size=8.2, bold=True, color="000000")
    query_run = paragraph.add_run(query)
    set_run_font(query_run, name="Consolas", size=7.7, color="222222")


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    for name, size in (("Title", 25), ("Heading 1", 16), ("Heading 2", 12.5)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(str(OUTPUT))
    clear_body(doc)
    configure_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(5)
    set_run_font(
        title.add_run("Informe técnico sobre brechas semánticas y trazabilidad en CIAR"),
        size=25,
        bold=True,
    )
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    set_run_font(
        subtitle.add_run("Evidencia desde los sílabos, la empleabilidad, el proceso de normalización y Neo4j"),
        size=11.5,
        italic=True,
        color="555555",
    )

    add_body(doc, "Fecha de corte: 5 de octubre de 2026 · Base consultada: Neo4j Desktop, instancia ulima_empleabilidad, base neo4j · Modalidad: consultas de solo lectura y revisión de salidas normalizadas.", space_after=10)

    add_heading(doc, "Resumen ejecutivo", 1)
    add_body(doc, "Este informe evalúa la calidad semántica, estructural y de trazabilidad de los datos que alimentan CIAR. El foco no es demostrar únicamente si una consulta Cypher se ejecuta, sino determinar si la información académica y laboral permite producir conclusiones semánticamente defendibles.")
    add_body(doc, "La evidencia confirma que una parte importante de las brechas se origina en los propios datos fuente. En la base actual, conceptos cercanos aparecen fragmentados entre distintos nombres, IDs, tipos y descripciones. El sistema puede almacenar y consultar correctamente esas entidades, pero la respuesta resultante puede subestimar, dividir o interpretar de forma incompleta la cobertura curricular y la demanda laboral.")
    add_body(doc, "La conclusión principal es que CIAR opera sobre una cadena de datos donde cada etapa puede conservar o amplificar una deficiencia anterior: sílabo u oferta, extracción, normalización, catálogo, publicación en Neo4j, consulta y redacción de la respuesta.")

    add_heading(doc, "1. Propósito y alcance", 1)
    add_body(doc, "El propósito es documentar con evidencia reproducible las brechas que afectan la integración entre información curricular e información de empleabilidad. El informe distingue problemas originados en la fuente, decisiones de normalización, limitaciones del modelo de datos y efectos sobre el agente.")
    add_body(doc, "El análisis utiliza el corte curricular 2026-2, la ejecución laboral de 2025 disponible localmente y consultas ejecutadas en la base Neo4j actualmente conectada. Las cifras pertenecen a cortes distintos y no deben interpretarse como un único universo estadístico.")
    add_body(doc, "Las herramientas se presentan como una dimensión histórica y parcialmente retirada. En el corte curricular revisado fueron omitidas por falta de evidencia suficiente; por tanto, el informe no las trata como una dimensión actualmente validada.")

    add_heading(doc, "2. Cadena de datos evaluada", 1)
    add_body(doc, "La evaluación sigue la siguiente cadena:")
    add_bullet(doc, "Fuente original: sílabos, logros, competencias, ofertas y requerimientos laborales.")
    add_bullet(doc, "Extracción: lectura de texto, campos estructurados y evidencias.")
    add_bullet(doc, "Normalización: asignación de nombres canónicos, IDs, tipos y relaciones.")
    add_bullet(doc, "Publicación: catálogos, coberturas y entidades almacenadas en Neo4j.")
    add_bullet(doc, "Consulta: recorrido del grafo y agregación de resultados.")
    add_bullet(doc, "Respuesta: interpretación redactada por el agente.")
    add_body(doc, "Esta cadena permite distinguir entre un dato ausente, un dato mal clasificado, dos datos equivalentes que permanecen separados y una respuesta que interpreta de forma excesiva un resultado técnicamente correcto.")

    add_heading(doc, "3. Cortes y evidencia disponible", 1)
    add_table(
        doc,
        ["Corte", "Registros principales", "Hallazgo relevante"],
        [
            ["Curricular 2026-2", "1,037 registros; 195 competencias; 4,814 logros; 7,811 relaciones", "Herramientas omitidas; no hay pendientes de decisión en la ejecución."],
            ["Empleabilidad 2025 local", "9,754 filas de origen; 3 ofertas procesadas; 13 requerimientos; 222 habilidades", "La ejecución local es una muestra de prueba, no un universo para porcentajes generales."],
            ["Neo4j actual", "76,523 nodos y 258,300 relaciones observadas en Desktop", "Existen Habilidad y Logros; la consulta a Herramienta devuelve que la etiqueta no existe."],
        ],
        [1.35, 2.15, 2.05],
    )
    add_body(doc, "La diferencia entre los cortes es parte de la evidencia: los archivos locales, las salidas de normalización y la base publicada no deben asumirse como idénticos sin verificar su versión y fecha de publicación.")

    add_heading(doc, "4. Taxonomía de brechas", 1)
    add_body(doc, "Para efectos de coordinación, cada brecha se analiza con cuatro preguntas: qué se observa en el dato, qué evidencia permite demostrarlo, qué efecto produce en el análisis y qué parte corresponde a la fuente, al proceso de normalización o al modelo publicado. Esta separación evita atribuir al agente una deficiencia que ya estaba presente en los sílabos, las ofertas o el catálogo de origen.")
    add_table(
        doc,
        ["Brecha", "Dónde se observa", "Consecuencia"],
        [
            ["Identidad", "Fuente y catálogo", "Conceptos similares reciben nombres o IDs distintos."],
            ["Duplicidad semántica", "Catálogo curricular", "Un mismo nombre representa descripciones diferentes."],
            ["Clasificación", "Extracción y normalización", "Una entidad puede cambiar entre genérica, específica, habilidad o logro."],
            ["Granularidad", "Integración", "Se comparan conceptos generales con conceptos específicos."],
            ["Terminología", "Currículo frente a mercado", "La misma capacidad puede expresarse con vocabularios distintos."],
            ["Completitud", "Fuente y publicación", "La ausencia textual puede confundirse con ausencia real."],
            ["Versionado", "Pipeline y esquema", "Un mismo campo o dimensión puede cambiar de significado."],
            ["Trazabilidad", "Neo4j y agente", "No siempre se puede reconstruir la evidencia que sustenta un resultado."],
        ],
        [1.25, 2.1, 2.2],
    )
    add_body(doc, "Los falsos positivos, falsos negativos y gaps aparentes no son brechas de origen independientes: son efectos que pueden producirse cuando se combinan las brechas anteriores.")

    add_heading(doc, "5. Evidencia de brechas en las competencias", 1)
    add_heading(doc, "5.1 Identidad y fragmentación", 2)
    add_body(doc, "La consulta siguiente recupera las competencias cuyos nombres contienen variantes de trabajo en equipo. La consulta devuelve 32 registros, lo que muestra que el concepto no está representado como una única entidad semántica en el grafo.")
    add_query(doc, "MATCH (c:Competencia)\nWHERE toLower(c.nombre_competencia) IN ['trabajo en equipo', 'trabajo colaborativo', 'trabajo cooperativo']\nRETURN c.nombre_competencia AS nombre, c.id_competencia AS id, c.tipo_competencia AS tipo, c.codigo_competencia AS codigo, coalesce(c.descripcion_breve_competencia, c.descripcion_breve) AS descripcion\nORDER BY nombre, id", "Consulta 1 · Variantes de una competencia")
    add_figure(doc, FIG_DIR / "evidencia_trabajo_equipo_neo4j.png", "Figura 1. Variantes e IDs de competencias relacionadas con trabajo en equipo en Neo4j")
    add_body(doc, "La agrupación por cobertura muestra una consecuencia cuantitativa de esta fragmentación:")
    add_table(
        doc,
        ["Nombre almacenado", "IDs", "Cursos", "Coberturas"],
        [
            ["Trabajo colaborativo", "26", "92", "278"],
            ["Trabajo Colaborativo", "4", "41", "137"],
            ["Trabajo cooperativo", "1", "1", "3"],
            ["Trabajo en equipo", "1", "1", "2"],
        ],
        [2.1, 0.8, 0.8, 1.0],
    )
    add_query(doc, "MATCH (c:Carrera)-[:ENSENIA]->(cu:Curso)-[:TIENE]->(cc:Cobertura_Curricular)-[:CUBRE]->(co:Competencia)\nWHERE toLower(co.nombre_competencia) IN ['trabajo en equipo', 'trabajo colaborativo', 'trabajo cooperativo']\nRETURN co.nombre_competencia AS competencia, count(DISTINCT co.id_competencia) AS ids, count(DISTINCT cu.id_curso) AS cursos, count(cc) AS coberturas\nORDER BY coberturas DESC", "Consulta 2 · Cobertura por variante")
    add_figure(doc, FIG_DIR / "evidencia_cobertura_trabajo_neo4j.png", "Figura 2. Cobertura separada por variante nominal en Neo4j")
    add_body(doc, "Este resultado no demuestra por sí solo que todas las variantes deban fusionarse. Sí demuestra que la base actual no puede tratarlas como una única familia conceptual sin una decisión adicional de equivalencia y sin revisar sus descripciones.")
    add_body(doc, "El mismo patrón aparece en otras competencias. En el catálogo curricular, Pensamiento crítico se registra en 44 filas, con variantes de mayúsculas y acentuación —Pensamiento crítico, Pensamiento Critico, Pensamiento Crítico y Pensamiento critico— y 10 descripciones distintas. El nombre parece equivalente, pero las descripciones pueden representar desempeños diferentes; por tanto, la normalización debe separar la corrección ortográfica de la decisión de equivalencia conceptual.")
    add_body(doc, "Otro ejemplo es Comportamiento ético: aparece en 29 registros y cinco descripciones. Aquí el problema no es solamente la escritura, sino que una misma etiqueta se utiliza para múltiples formulaciones de resultado. Consolidar por nombre sin revisar la descripción podría sobreestimar la cobertura de un desempeño específico.")
    add_figure(doc, FIG_DIR / "evidencia_catalogo_competencias_variantes.png", "Figura 3. Evidencia exportada de variantes, IDs y descripciones en el catálogo de competencias")

    add_heading(doc, "5.2 Duplicidad de nombre y diferencia de descripción", 2)
    add_body(doc, "La misma etiqueta puede representar contenidos distintos. Por ejemplo, Trabajo colaborativo aparece asociado a descripciones sobre equipos multidisciplinarios, diálogo, comercio electrónico, desarrollo de software y ambientes inclusivos. La duplicidad no es solamente ortográfica: también afecta el significado de la entidad.")
    add_body(doc, "La consecuencia es que un conteo por nombre puede mezclar conceptos distintos, mientras que un conteo por ID puede fragmentar conceptos equivalentes. Ambos resultados requieren interpretación y evidencia de origen.")
    add_body(doc, "En el caso de Trabajo colaborativo, el catálogo presenta 28 registros, ocho descripciones y dos tipos de competencia: genérica y específica. Una consulta que filtre solamente por nombre recuperará ambos niveles; una consulta que filtre por tipo puede excluir parte de la cobertura; y una consulta que cuente IDs puede presentar 28 entidades cuando la coordinación está preguntando por una sola familia de capacidad. Cada lectura responde una pregunta distinta.")
    add_body(doc, "El riesgo concreto es doble. Si se fusiona todo, se puede atribuir a una competencia genérica una cobertura que en realidad corresponde a una aplicación disciplinar. Si no se fusiona nada, se puede informar que la capacidad aparece menos veces de lo que aparece en los sílabos. La decisión debe quedar registrada como equivalencia, relación o distinción.")

    add_heading(doc, "5.3 Clasificación y granularidad", 2)
    add_body(doc, "En el conjunto revisado, Trabajo colaborativo aparece principalmente como competencia genérica, pero también como competencia específica. Esta diferencia puede alterar los análisis que separan competencias transversales de competencias disciplinares.")
    add_body(doc, "También se observan diferencias de nivel entre competencia, logro, habilidad y herramienta. Una mención a Python en un logro curricular no equivale automáticamente a una herramienta publicada como entidad independiente, y una competencia general de Inteligencia Artificial no prueba cobertura específica de Machine Learning o Deep Learning.")
    add_table(
        doc,
        ["Caso", "Evidencia observada", "Brecha que produce"],
        [
            ["Trabajo colaborativo", "28 registros; tipos genérico y específico; 8 descripciones", "El nombre no determina el nivel ni el desempeño."],
            ["Pensamiento crítico", "44 registros; 10 descripciones y variantes ortográficas", "El conteo nominal mezcla formulaciones distintas."],
            ["Liderazgo", "5 registros; 2 descripciones", "Un mismo nombre puede representar desempeños diferentes."],
            ["IA, Machine Learning y Deep Learning", "Relación de general a específica no demostrada por el nombre", "La cobertura general no prueba dominio técnico específico."],
        ],
        [1.45, 2.7, 1.75],
    )
    add_body(doc, "La granularidad también afecta la explicación de una ausencia. Si una oferta solicita una habilidad técnica específica y el currículo solo contiene una competencia general, el resultado correcto no es afirmar automáticamente que la habilidad no existe; debe clasificarse como cobertura no demostrada o cobertura pendiente de equivalencia, según la evidencia textual.")
    add_heading(doc, "5.4 Lectura correcta de la responsabilidad", 2)
    add_body(doc, "En este caso, la brecha semántica es compartida. Los sílabos pueden declarar la misma capacidad con nombres distintos —por ejemplo, Trabajo en equipo y Trabajo colaborativo— y el proceso puede conservar esa diferencia para no fusionar datos sin evidencia. La responsabilidad del sistema es hacer visible la equivalencia posible, conservar los IDs originales y advertir que el conteo por etiqueta no representa necesariamente el concepto consolidado.")
    add_body(doc, "Por ello, la mejora no consiste únicamente en modificar la consulta. Requiere una decisión de catálogo: definir si las variantes son equivalentes, relacionadas o distintas; registrar la decisión; y volver a calcular la cobertura con la misma regla en todos los periodos.")
    add_heading(doc, "5.5 Normalización de atributos descriptivos", 2)
    add_body(doc, "No todas las brechas impactan directamente la semántica de una competencia. Algunas afectan filtros, agrupaciones y reportes operativos. En el catálogo de cursos, el mismo valor aparece con distintas mayúsculas, acentos, guiones y espacios.")
    add_table(
        doc,
        ["Atributo", "Variantes observadas", "Registros afectados", "Efecto"],
        [
            ["tipo_curso", "Obligatoria / OBLIGATORIA", "475", "Agrupaciones por texto pueden dividir el mismo tipo."],
            ["creditos", "Tres (3) / TRES (3)", "550", "Los conteos por créditos pueden duplicar categorías."],
            ["naturaleza", "Múltiples formas de Teórico-práctica", "546", "Los filtros de naturaleza dependen de la normalización."],
            ["nivel", "Décimo / Decimo", "77", "La ausencia de tilde puede crear una categoría adicional."],
        ],
        [1.2, 2.1, 1.0, 1.7],
    )
    add_body(doc, "Este caso debe clasificarse como brecha de normalización de atributos, no como duplicidad curricular de cursos. La identidad del curso está respaldada por un ID; sin embargo, los análisis que agrupen por el valor visible necesitan una representación canónica o una función de normalización explícita.")
    add_figure(doc, FIG_DIR / "evidencia_catalogo_cursos_normalizacion.png", "Figura 4. Evidencia exportada de variantes en atributos de curso")
    add_heading(doc, "5.6 Duplicidad de habilidades técnicas por carrera", 2)
    add_body(doc, "El catálogo de habilidades mantiene el ámbito de carrera en el identificador, lo cual evita fusionar automáticamente habilidades que podrían tener distinto contexto profesional. A la vez, esta decisión produce un posible gap de identidad entre carreras.")
    add_table(
        doc,
        ["Nombre de habilidad", "IDs encontrados", "Lectura de la brecha"],
        [
            ["Desarrollar métodos o procesos técnicos.", "4", "Misma formulación repetida en cuatro carreras; conviene determinar si es la misma habilidad transversal o cuatro instancias contextualizadas."],
            ["Analizar datos empresariales o financieros.", "2", "El match por ID puede separar una capacidad equivalente presente en más de una carrera."],
            ["Implementar mejoras de diseño o de procesos.", "3", "La cobertura puede parecer fragmentada aunque la formulación sea equivalente."],
            ["Preparar documentos de procedimiento.", "3", "El nombre común no garantiza que el desempeño sea idéntico en cada carrera."],
        ],
        [2.0, 0.8, 3.2],
    )
    add_body(doc, "La decisión técnica de mantener IDs por carrera es defendible para preservar contexto y trazabilidad. La brecha aparece cuando se interpreta el ID como si fuera la única identidad semántica: puede producir falsos negativos entre carreras o impedir que la coordinación vea una capacidad transversal. La solución requiere una clave canónica de habilidad y, cuando corresponda, una relación de equivalencia o especialización entre IDs.")
    add_figure(doc, FIG_DIR / "evidencia_catalogo_habilidades_ids.png", "Figura 5. Evidencia exportada de habilidades repetidas en distintos IDs y carreras")

    add_heading(doc, "6. Evidencia de integración curricular y laboral", 1)
    add_body(doc, "La integración más fuerte disponible en el grafo ocurre cuando una misma entidad canónica Habilidad aparece conectada tanto a una cobertura curricular como a un requerimiento laboral. Este tipo de coincidencia es verificable, pero no debe confundirse con una equivalencia semántica completa: la identidad canónica demuestra una coincidencia del modelo, no necesariamente la suficiencia profesional de la cobertura.")
    add_query(doc, "MATCH (s:Silabo)-[:DECLARA]->(cc:Cobertura_Curricular)-[:CUBRE]->(h:Habilidad)\nMATCH (o:Oferta_Laboral)-[:TIENE]->(r:Requerimiento_Laboral)-[:REQUIERE]->(h)\nRETURN h.nombre_habilidad AS habilidad, count(DISTINCT o) AS ofertas, count(DISTINCT s) AS silabos\nORDER BY ofertas DESC, habilidad\nLIMIT 20", "Consulta 3 · Intersección curricular y laboral")
    add_figure(doc, FIG_DIR / "evidencia_interseccion_curricular_laboral.png", "Figura 6. Habilidades compartidas por ofertas laborales y sílabos mediante la misma entidad canónica")
    add_body(doc, "Los primeros resultados muestran habilidades con alta presencia laboral y presencia curricular en un número reducido de sílabos. Esto puede significar una cobertura curricular concentrada, una diferencia de población entre los cortes o una publicación curricular incompleta. No debe interpretarse automáticamente como una brecha real sin revisar el universo, la carrera, el periodo y la evidencia textual.")

    add_heading(doc, "7. Versionado y dimensión de herramientas", 1)
    add_body(doc, "Las herramientas deben documentarse como una dimensión de evolución del sistema. En versiones anteriores se trabajó con catálogos y relaciones de herramientas. Posteriormente, el pipeline curricular cambió el contrato de habilidades hacia logros y algunos cortes omitieron explícitamente las herramientas por falta de datos suficientes.")
    add_query(doc, "MATCH (n:Habilidad) RETURN 'Habilidad' AS etiqueta, count(n) AS total\nUNION ALL\nMATCH (n:Logros) RETURN 'Logros' AS etiqueta, count(n) AS total\nUNION ALL\nMATCH (n:Herramienta) RETURN 'Herramienta' AS etiqueta, count(n) AS total\nORDER BY etiqueta", "Consulta 4 · Etiquetas activas y etiqueta ausente")
    add_figure(doc, FIG_DIR / "evidencia_etiquetas_habilidad_logros.png", "Figura 7. Etiquetas activas en Neo4j y ausencia de la etiqueta Herramienta en la consulta actual")
    add_body(doc, "La consulta devuelve 1,792 nodos Habilidad y 4,366 nodos Logros. La consulta a Herramienta devuelve el aviso de que la etiqueta no existe. Esto debe reportarse como una limitación actual del corte, no como una falla de la consulta ni como evidencia de que la dimensión de herramientas esté validada.")

    add_heading(doc, "8. Propagación de las brechas hacia el análisis", 1)
    add_body(doc, "Las brechas anteriores pueden producir cuatro efectos principales:")
    add_bullet(doc, "Falso positivo: se declara cobertura porque dos conceptos son relacionados, aunque no equivalentes.")
    add_bullet(doc, "Falso negativo: se declara ausencia porque los nombres no coinciden, aunque la capacidad esté presente.")
    add_bullet(doc, "Gap aparente: la diferencia proviene de terminología, clasificación o granularidad.")
    add_bullet(doc, "Gap indeterminado: no existe evidencia suficiente para decidir debido a datos omitidos o muestras incompletas.")
    add_body(doc, "El análisis de brechas debe incorporar el tamaño y la fecha del universo. La ejecución laboral local de 2025 contiene 9,754 filas de origen, pero solo tres ofertas procesadas en la ejecución de prueba. Por ello no es válido usar esa ejecución para concluir porcentajes generales del mercado.")
    add_table(
        doc,
        ["Tipo de conclusión", "Condición mínima", "Tratamiento en el informe"],
        [
            ["Brecha real", "Demanda comprobada y ausencia curricular comprobada", "Presentar con evidencia de ambos dominios."],
            ["Brecha aparente", "Nombres o niveles diferentes", "Presentar como hipótesis pendiente de equivalencia."],
            ["Cobertura implícita", "Evidencia textual indirecta", "No convertir automáticamente en cobertura confirmada."],
            ["Brecha indeterminada", "Datos faltantes o corte incompleto", "Informar la limitación y no concluir ausencia."],
        ],
        [1.25, 2.2, 2.2],
    )

    add_heading(doc, "9. Trazabilidad y auditabilidad", 1)
    add_body(doc, "Una respuesta del agente debe poder reconstruirse desde el resultado hasta la evidencia original. Para cada afirmación agregada, la trazabilidad ideal incluye oferta, requerimiento, habilidad, sílabo, curso, carrera, periodo, ID de ejecución y evidencia textual.")
    add_body(doc, "La trazabilidad también permite distinguir un problema del dato de un problema del agente. Si Neo4j contiene dos entidades separadas, la consulta puede estar correcta y aun así producir un resultado semánticamente fragmentado. El informe debe mostrar ambos niveles de evaluación.")

    add_heading(doc, "10. Evaluación del agente", 1)
    add_body(doc, "La evaluación debe combinar calidad técnica y calidad semántica. No basta con que Cypher se ejecute sin error.")
    add_table(
        doc,
        ["Dimensión", "Qué se verifica"],
        [
            ["Consulta", "Solo lectura, labels y relaciones válidos, parámetros correctos."],
            ["Resultado", "Filas, conteos, orden y filtros correctos."],
            ["Semántica", "La interpretación no confunde identidad, relación o cobertura."],
            ["Trazabilidad", "La respuesta puede regresar a sus fuentes."],
            ["Limitaciones", "La respuesta advierte sobre muestras, versiones y datos faltantes."],
        ],
        [1.35, 4.2],
    )
    add_body(doc, "Las preguntas de aceptación deben cubrir consultas curriculares, laborales, intersecciones y brechas. Cada una debe tener una respuesta esperada revisada manualmente y una clasificación de posibles errores.")

    add_heading(doc, "11. Match entre cobertura curricular y oferta laboral mediante habilidades técnicas", 1)
    add_body(doc, "El match analizado en este apartado se define de manera estricta: existe coincidencia cuando la misma entidad canónica Habilidad está conectada, por un lado, con una cobertura curricular y, por otro, con un requerimiento laboral. El match no se calcula únicamente por similitud textual y no demuestra por sí solo que la cobertura sea suficiente, reciente o distribuida entre varias carreras.")
    add_body(doc, "El archivo CSV descargado el 5 de octubre de 2026 contiene 134 habilidades técnicas con coincidencia en ambos dominios. En conjunto registra 171 coberturas curriculares y 34,293 requerimientos laborales. La cobertura es positiva en las 134 filas, pero está concentrada: 103 habilidades aparecen con una sola cobertura curricular; de ellas, 58 superan 100 requerimientos laborales. Esta concentración es el hallazgo principal para coordinación: que exista match no significa que la capacidad esté integrada de forma amplia o robusta en el currículo.")
    add_table(
        doc,
        ["Habilidad técnica", "Coberturas", "Requerimientos", "Curso asociado", "Sílabo"],
        [
            ["Analizar datos para informar decisiones o actividades operativas.", "4", "1,539", "Estadística Básica para los Negocios", "6502"],
            ["Preparar informes operativos.", "1", "1,381", "Introducción a la Ingeniería", "510005"],
            ["Analizar datos empresariales o financieros.", "2", "1,072", "Herramientas Informáticas para la Gestión I", "520004"],
            ["Dirigir actividades de ventas, marketing o servicio al cliente.", "2", "993", "Gestión Comercial", "520050"],
            ["Analizar información financiera.", "1", "858", "Fundamentos de Finanzas", "540047"],
        ],
        [2.7, 0.65, 0.85, 1.65, 0.55],
    )
    add_figure(doc, FIG_DIR / "evidencia_match_csv_habilidades_tecnicas.png", "Figura 8. Evidencia exportada de habilidades técnicas con match curricular y laboral")
    add_body(doc, "Los ejemplos muestran tres situaciones distintas. Primero, una habilidad con demanda muy alta y cuatro coberturas, que puede representar una base curricular relativamente más distribuida. Segundo, habilidades con más de mil requerimientos y una sola cobertura, que deben revisarse como posibles casos de concentración o de cobertura insuficientemente distribuida. Tercero, una coincidencia con curso y sílabo identificables, que permite iniciar la trazabilidad hasta la fuente académica y revisar si el texto del sílabo realmente evidencia el desempeño esperado.")
    add_body(doc, "El match también debe revisarse a nivel de identidad. Analizar datos empresariales o financieros aparece en el catálogo curricular con dos IDs, mientras que Desarrollar métodos o procesos técnicos aparece con cuatro IDs asociados a distintas carreras. Si la oferta laboral utiliza una de esas identidades y el currículo utiliza otra, el cruce exacto puede no encontrar la relación aunque el texto sea equivalente. En sentido contrario, fusionar ambos IDs sin revisar la descripción puede ocultar diferencias disciplinares. Por eso, el indicador debe reportar coincidencia exacta, equivalencia validada y casos no resueltos por separado.")
    add_body(doc, "La lectura de match recomendada para coordinación es la siguiente: (1) confirmar la identidad canónica de la habilidad; (2) revisar la cantidad de requerimientos laborales y su distribución; (3) identificar cursos, sílabos, carreras y periodos que sostienen la cobertura; (4) revisar la evidencia textual; y (5) clasificar el resultado como cobertura suficiente, cobertura concentrada, cobertura implícita o caso pendiente. El informe no debe presentar el número de requerimientos como porcentaje de empleabilidad ni como indicador causal de pertinencia curricular.")
    add_body(doc, "El CSV también permite observar que la salida actual está construida sobre intersecciones ya encontradas. Por tanto, sus 134 filas no representan todas las habilidades solicitadas por el mercado ni todas las habilidades declaradas en los sílabos; representan el subconjunto que logró aparecer en ambos lados con la regla de coincidencia vigente.")

    add_heading(doc, "12. Conclusiones", 1)
    add_body(doc, "La evidencia actual confirma que las brechas no pertenecen exclusivamente al agente. Una parte se origina en la redacción de los sílabos y las ofertas; otra se conserva durante la normalización para evitar fusiones no justificadas; y otra se manifiesta cuando el modelo de datos intenta representar conceptos de diferente tipo o nivel.")
    add_body(doc, "El caso de Trabajo en equipo, Trabajo colaborativo y Trabajo cooperativo demuestra que el grafo puede estar estructuralmente válido y, aun así, fragmentar una familia conceptual. El caso de las herramientas demuestra que el versionado y la disponibilidad de datos deben formar parte explícita de la interpretación. La intersección entre Habilidad curricular y laboral demuestra que una coincidencia canónica es útil, pero no sustituye la evaluación de suficiencia semántica.")
    add_body(doc, "Los catálogos revisados muestran además que la brecha se extiende a atributos operativos y a habilidades repetidas por carrera. Las variantes de tipo de curso, créditos, naturaleza y nivel pueden afectar filtros y conteos, mientras que la repetición de una habilidad con varios IDs puede producir falsos negativos en el match exacto. Estos casos no deben corregirse con fusiones automáticas: requieren reglas de normalización, equivalencias documentadas y conservación de la procedencia.")
    add_body(doc, "Por tanto, las respuestas de CIAR deben distinguir entre cobertura confirmada, relación posible, ausencia de evidencia y brecha real. Esa distinción es necesaria para que el análisis sea reproducible, auditable y útil para la toma de decisiones académicas.")

    add_heading(doc, "13. Recomendaciones", 1)
    add_bullet(doc, "Mantener un catálogo institucional de competencias con IDs y descripciones estables.")
    add_bullet(doc, "Registrar variantes y posibles equivalencias sin fusionarlas automáticamente.")
    add_bullet(doc, "Crear una clave canónica transversal para habilidades repetidas por carrera y mantener relaciones de equivalencia o especialización hacia los IDs contextuales.")
    add_bullet(doc, "Normalizar atributos de curso —tipo, créditos, naturaleza y nivel— conservando el valor original para auditoría.")
    add_bullet(doc, "Conservar la procedencia hasta el texto del sílabo u oferta que sustenta cada relación.")
    add_bullet(doc, "Separar explícitamente competencia, habilidad, logro y herramienta en los contratos de datos.")
    add_bullet(doc, "Reportar siempre el corte, periodo, tamaño de muestra y componentes omitidos.")
    add_bullet(doc, "Incluir una advertencia cuando la conclusión dependa de cobertura implícita o equivalencia no validada.")
    add_bullet(doc, "Retomar la dimensión de herramientas únicamente cuando exista evidencia suficiente y un contrato estable.")

    add_heading(doc, "Anexo A. Consultas ejecutadas", 1)
    add_body(doc, "Las consultas incluidas en este anexo fueron ejecutadas en Neo4j Desktop con acceso de lectura. Se conservaron capturas de los resultados que sustentan los casos principales.")
    add_query(doc, "MATCH (c:Competencia)\nWHERE toLower(c.nombre_competencia) IN ['trabajo en equipo', 'trabajo colaborativo', 'trabajo cooperativo']\nRETURN c.nombre_competencia AS nombre, c.id_competencia AS id, c.tipo_competencia AS tipo, c.codigo_competencia AS codigo, coalesce(c.descripcion_breve_competencia, c.descripcion_breve) AS descripcion\nORDER BY nombre, id", "A.1 Variantes de competencia")
    add_query(doc, "MATCH (c:Carrera)-[:ENSENIA]->(cu:Curso)-[:TIENE]->(cc:Cobertura_Curricular)-[:CUBRE]->(co:Competencia)\nWHERE toLower(co.nombre_competencia) IN ['trabajo en equipo', 'trabajo colaborativo', 'trabajo cooperativo']\nRETURN co.nombre_competencia AS competencia, count(DISTINCT co.id_competencia) AS ids, count(DISTINCT cu.id_curso) AS cursos, count(cc) AS coberturas\nORDER BY coberturas DESC", "A.2 Cobertura por variante")
    add_query(doc, "MATCH (s:Silabo)-[:DECLARA]->(cc:Cobertura_Curricular)-[:CUBRE]->(h:Habilidad)\nMATCH (o:Oferta_Laboral)-[:TIENE]->(r:Requerimiento_Laboral)-[:REQUIERE]->(h)\nRETURN h.nombre_habilidad AS habilidad, count(DISTINCT o) AS ofertas, count(DISTINCT s) AS silabos\nORDER BY ofertas DESC, habilidad\nLIMIT 20", "A.3 Intersección curricular y laboral")
    add_query(doc, "MATCH (n:Habilidad) RETURN 'Habilidad' AS etiqueta, count(n) AS total\nUNION ALL\nMATCH (n:Logros) RETURN 'Logros' AS etiqueta, count(n) AS total\nUNION ALL\nMATCH (n:Herramienta) RETURN 'Herramienta' AS etiqueta, count(n) AS total\nORDER BY etiqueta", "A.4 Disponibilidad de etiquetas")

    add_heading(doc, "Anexo B. Fuentes de evidencia", 1)
    add_bullet(doc, "Segundo Informe.docx: documento de referencia conceptual y de diseño.")
    add_bullet(doc, "backend/agente/utils/schema_ciar.py: contrato estático de la ontología del agente.")
    add_bullet(doc, "backend/agente/dashboard/consultas.py: consultas de demanda, cobertura y brechas.")
    add_bullet(doc, "backend/.normalizador/NOR_dc327a683a8d4933/manifest.json: ejecución curricular 2026-2.")
    add_bullet(doc, "backend/.normalizador/NOR_52a595f2e7c54138/manifest.json: ejecución laboral 2025 de prueba.")
    add_bullet(doc, "Catálogos descargados del 5 de octubre de 2026: competencias, logros, habilidades, coberturas, cursos y sílabos.")
    add_bullet(doc, "Capturas de Neo4j Desktop incluidas en este informe.")

    add_heading(doc, "Anexo C. Archivo CSV descargado y soporte del match", 1)
    add_body(doc, "Archivo analizado: neo4j_query_table_data_2026-10-5.csv. Ubicación de referencia: C:\\Users\\I13305\\Downloads. Fecha de modificación observada: 5 de octubre de 2026, 11:38. Tamaño observado: 93,692 bytes. El archivo contiene 134 registros de habilidades técnicas y las siguientes columnas: id_habilidad, habilidad, coberturas_curriculares, requerimientos_laborales, cursos, silabos, cargos y empresas.")
    add_body(doc, "El archivo se utilizó como soporte del apartado 11, sin modificar su contenido. La validación de lectura comprobó que las 134 filas tienen valores no vacíos en las ocho columnas; las coberturas curriculares oscilan entre 1 y 4, y los requerimientos laborales entre 1 y 1,539 por habilidad. Los nombres de cursos, cargos y empresas se conservan como listas exportadas y deben interpretarse como evidencia de trazabilidad, no como conteos adicionales.")
    add_table(
        doc,
        ["Campo", "Uso en el análisis", "Observación"],
        [
            ["id_habilidad", "Identidad canónica", "Permite evitar que el nombre sea la única clave."],
            ["habilidad", "Descripción legible", "Debe revisarse junto con el ID y la evidencia textual."],
            ["coberturas_curriculares", "Presencia curricular", "No equivale a número de carreras ni suficiencia."],
            ["requerimientos_laborales", "Demanda observada", "No equivale a número de ofertas únicas."],
            ["cursos y silabos", "Trazabilidad académica", "Permiten ubicar el origen de la cobertura."],
            ["cargos y empresas", "Contexto laboral", "Representan la lista exportada asociada a la habilidad."],
        ],
        [1.25, 2.1, 2.5],
    )
    add_body(doc, "Muestra de registros priorizados por requerimientos laborales:")
    add_table(
        doc,
        ["ID", "Habilidad", "Coberturas", "Requerimientos"],
        [
            ["HAB_TEC_0009", "Analizar datos para informar decisiones o actividades operativas.", "4", "1,539"],
            ["HAB_TEC_0159", "Preparar informes operativos.", "1", "1,381"],
            ["HAB_TEC_0007", "Analizar datos empresariales o financieros.", "2", "1,072"],
            ["HAB_TEC_0077", "Dirigir actividades de ventas, marketing o servicio al cliente.", "2", "993"],
            ["HAB_TEC_0014", "Analizar información financiera.", "1", "858"],
        ],
        [1.0, 3.6, 0.8, 1.0],
    )
    add_body(doc, "El CSV debe conservarse junto con el informe y su fecha de corte. Si se vuelve a generar, los resultados deben compararse por id_habilidad, corte, regla de match y versión del catálogo; de lo contrario, una variación en el número de coincidencias podría atribuirse erróneamente a una mejora curricular cuando en realidad proviene de un cambio en los datos o en la normalización.")

    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build()
