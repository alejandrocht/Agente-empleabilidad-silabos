"""Catálogos versionados de habilidades técnicas e índices locales de embeddings.

El Excel es la fuente de verdad. El índice binario es un artefacto derivado y
reproducible: nunca se usa para asignar ni modificar IDs de HAB_TEC.
"""

from __future__ import annotations

import hashlib
import json
import math
import re
import shutil
import struct
import threading
import unicodedata
import urllib.error
import urllib.request
from collections.abc import Callable, Iterable
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from agente.config.settings import BASE_DIR, entero, texto
from agente.normalizador.hab_tec_neo4j import publicador_hab_tec_neo4j


def normalize(value: object) -> str:
    """Normaliza etiquetas de carrera sin depender del módulo legado eliminado."""

    texto = unicodedata.normalize("NFKD", str(value or ""))
    texto = "".join(caracter for caracter in texto if not unicodedata.combining(caracter))
    return " ".join(texto.casefold().split())

_COLUMNAS = ("carrera", "id", "nombre", "descripcion")
_ESTADOS_FINALES = frozenset({"vectorizado", "rechazado", "error"})


class ErrorCatalogoHabTec(ValueError):
    """Error seguro, apropiado para mostrar en la interfaz administrativa."""


@dataclass(frozen=True, slots=True)
class HallazgoCatalogo:
    codigo: str
    severidad: str
    mensaje: str
    fila: int | None = None
    detalle: str = ""

    def a_dict(self) -> dict[str, object]:
        return {
            "codigo": self.codigo,
            "severidad": self.severidad,
            "mensaje": self.mensaje,
            "fila": self.fila,
            "detalle": self.detalle,
        }


@dataclass(frozen=True, slots=True)
class HabilidadTecnica:
    carrera: str
    id_hab_tec: str
    nombre: str
    descripcion: str
    fila: int

    @property
    def texto_indexado(self) -> str:
        return (
            f"Habilidad técnica: {self.nombre}.\n"
            f"Descripción: {self.descripcion}.\n"
            f"Carrera: {self.carrera}."
        )

    def a_dict(self, *, incluir_texto: bool = False) -> dict[str, object]:
        resultado: dict[str, object] = {
            "carrera": self.carrera,
            "id": self.id_hab_tec,
            "nombre": self.nombre,
            "descripcion": self.descripcion,
            "fila": self.fila,
        }
        if incluir_texto:
            resultado["texto_indexado"] = self.texto_indexado
        return resultado


@dataclass(slots=True)
class CatalogoHabTec:
    id_catalogo: str
    archivo: str
    directorio: Path
    creado_en: str
    estado: str
    hoja: str = ""
    registros: tuple[HabilidadTecnica, ...] = ()
    hallazgos: tuple[HallazgoCatalogo, ...] = ()
    modelo_embedding: str = ""
    dimension_embedding: int = 0
    resumen_neo4j: dict[str, int] | None = None
    error: str = ""
    actualizada_en: str = ""

    def valida(self) -> bool:
        return not any(item.severidad == "error" for item in self.hallazgos)

    def a_dict(self) -> dict[str, object]:
        errores = sum(item.severidad == "error" for item in self.hallazgos)
        advertencias = sum(item.severidad == "warning" for item in self.hallazgos)
        return {
            "id_catalogo": self.id_catalogo,
            "archivo": self.archivo,
            "estado": self.estado,
            "creado_en": self.creado_en,
            "actualizada_en": self.actualizada_en or self.creado_en,
            "hoja": self.hoja,
            "valido": self.valida(),
            "resumen": {
                "habilidades": len(self.registros),
                "errores": errores,
                "advertencias": advertencias,
                "modelo_embedding": self.modelo_embedding or None,
                "dimension_embedding": self.dimension_embedding or None,
                "neo4j": self.resumen_neo4j,
            },
            "preview": [item.a_dict() for item in self.registros[:100]],
            "hallazgos": [item.a_dict() for item in self.hallazgos[:100]],
            "error": self.error or None,
        }


def _ahora() -> str:
    return datetime.now(UTC).isoformat()


def _texto(valor: object) -> str:
    return re.sub(r"\s+", " ", str(valor or "")).strip()


def _clave_columna(valor: object) -> str:
    return normalize(valor).replace(" ", "")


def _huella(ruta: Path) -> str:
    digest = hashlib.sha256()
    with ruta.open("rb") as archivo:
        for bloque in iter(lambda: archivo.read(1024 * 1024), b""):
            digest.update(bloque)
    return digest.hexdigest()


def leer_catalogo_excel(
    ruta: Path,
) -> tuple[str, tuple[HabilidadTecnica, ...], tuple[HallazgoCatalogo, ...]]:
    """Lee el primer sheet que contenga el contrato de cuatro columnas."""

    if ruta.suffix.lower() != ".xlsx":
        raise ErrorCatalogoHabTec("El catálogo debe cargarse como archivo XLSX.")
    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except Exception as exc:
        raise ErrorCatalogoHabTec("No se pudo abrir el archivo XLSX.") from exc

    try:
        seleccion = None
        for nombre in libro.sheetnames:
            hoja = libro[nombre]
            cabecera = next(hoja.iter_rows(min_row=1, max_row=1, values_only=True), ())
            indices = {_clave_columna(valor): posicion for posicion, valor in enumerate(cabecera)}
            if all(columna in indices for columna in _COLUMNAS):
                seleccion = hoja, indices
                break
        if seleccion is None:
            raise ErrorCatalogoHabTec(
                "No se encontró una hoja con las columnas Carrera, id, nombre y descripcion."
            )

        hoja, indices = seleccion
        hallazgos: list[HallazgoCatalogo] = []
        registros: list[HabilidadTecnica] = []
        # Un HAB_TEC puede estar asociado a varias carreras. Por eso la clave
        # de una fila es la pareja Carrera + id y no el id aislado: cada fila
        # conserva su descripción contextual para la recuperación semántica.
        asociaciones: dict[tuple[str, str], int] = {}
        nombres: dict[str, tuple[str, int]] = {}
        ids: dict[str, tuple[str, int]] = {}
        for numero, valores in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
            if not any(_texto(valor) for valor in valores):
                continue
            fila = {
                columna: _texto(valores[indices[columna]])
                if indices[columna] < len(valores)
                else ""
                for columna in _COLUMNAS
            }
            faltantes = [columna for columna, valor in fila.items() if not valor]
            if faltantes:
                hallazgos.append(
                    HallazgoCatalogo(
                        "HAB_TEC_CAMPO_OBLIGATORIO_AUSENTE",
                        "error",
                        "La fila no contiene todos los campos obligatorios.",
                        numero,
                        ", ".join(faltantes),
                    )
                )
                continue
            clave_asociacion = (normalize(fila["carrera"]), normalize(fila["id"]))
            if clave_asociacion in asociaciones:
                hallazgos.append(
                    HallazgoCatalogo(
                        "HAB_TEC_ASOCIACION_DUPLICADA",
                        "error",
                        "La combinación de Carrera e ID de HAB_TEC debe ser única.",
                        numero,
                        f"También aparece en la fila {asociaciones[clave_asociacion]}.",
                    )
                )
                continue
            asociaciones[clave_asociacion] = numero
            clave_nombre = normalize(fila["nombre"])
            id_normalizado = normalize(fila["id"])
            nombre_previo_por_id = ids.get(id_normalizado)
            if nombre_previo_por_id and nombre_previo_por_id[0] != clave_nombre:
                hallazgos.append(
                    HallazgoCatalogo(
                        "HAB_TEC_ID_NOMBRE_VARIANTE",
                        "warning",
                        "Un mismo ID de HAB_TEC tiene nombres distintos; "
                        "se usará el primero como canónico.",
                        numero,
                        f"También aparece en la fila {nombre_previo_por_id[1]}.",
                    )
                )
            if not nombre_previo_por_id:
                ids[id_normalizado] = (clave_nombre, numero)
            nombre_previo = nombres.get(clave_nombre)
            if nombre_previo and nombre_previo[0] != id_normalizado:
                hallazgos.append(
                    HallazgoCatalogo(
                        "HAB_TEC_NOMBRE_REPETIDO",
                        "warning",
                        "El nombre aparece más de una vez; se conservarán IDs separados.",
                        numero,
                        f"También aparece en la fila {nombre_previo[1]}.",
                    )
                )
            elif not nombre_previo:
                nombres[clave_nombre] = (id_normalizado, numero)
            registros.append(
                HabilidadTecnica(
                    carrera=fila["carrera"],
                    id_hab_tec=fila["id"],
                    nombre=fila["nombre"],
                    descripcion=fila["descripcion"],
                    fila=numero,
                )
            )
        if not registros:
            hallazgos.append(
                HallazgoCatalogo(
                    "HAB_TEC_CATALOGO_VACIO",
                    "error",
                    "No se encontraron filas válidas de HAB_TEC.",
                )
            )
        return str(hoja.title), tuple(registros), tuple(hallazgos)
    finally:
        libro.close()


class GestorCatalogosHabTec:
    """Persiste catálogos e indexa con Ollama sin depender de una base vectorial."""

    def __init__(
        self,
        base_dir: Path | None = None,
        publicar_en_neo4j: (
            Callable[[Iterable[HabilidadTecnica], str], dict[str, int]] | None
        ) = None,
    ) -> None:
        configurado = texto("NORMALIZADOR_HAB_TEC_DIR")
        self.base_dir = base_dir or Path(configurado or BASE_DIR / ".normalizador" / "hab_tec")
        self._catalogos: dict[str, CatalogoHabTec] = {}
        self._bloqueo = threading.Lock()
        self._executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="hab-tec")
        self._publicar_en_neo4j = publicar_en_neo4j or publicador_hab_tec_neo4j.publicar

    def crear_desde_archivo(self, archivo: str, ruta_temporal: Path) -> dict[str, object]:
        identificador = f"HABTEC_{uuid4().hex[:16]}"
        directorio = self.base_dir / identificador
        directorio.mkdir(parents=True, exist_ok=False)
        nombre = Path(archivo).name or "catalogo_hab_tec.xlsx"
        entrada = directorio / "catalogo_hab_tec.xlsx"
        shutil.move(str(ruta_temporal), entrada)
        creado = _ahora()
        try:
            hoja, registros, hallazgos = leer_catalogo_excel(entrada)
        except ErrorCatalogoHabTec as exc:
            catalogo = CatalogoHabTec(
                identificador,
                nombre,
                directorio,
                creado,
                "rechazado",
                error=str(exc),
                actualizada_en=_ahora(),
            )
        else:
            estado = "listo_para_vectorizar"
            if any(item.severidad == "error" for item in hallazgos):
                estado = "rechazado"
            catalogo = CatalogoHabTec(
                identificador,
                nombre,
                directorio,
                creado,
                estado,
                hoja,
                registros,
                hallazgos,
                actualizada_en=_ahora(),
            )
            self._escribir_registros(catalogo)
        with self._bloqueo:
            self._catalogos[identificador] = catalogo
        self._persistir(catalogo)
        return catalogo.a_dict()

    def vectorizar(self, identificador: str) -> dict[str, object]:
        catalogo = self._obtener(identificador)
        if catalogo.estado != "listo_para_vectorizar":
            raise ErrorCatalogoHabTec("El catálogo no está disponible para vectorizar.")
        catalogo.estado = "vectorizando"
        catalogo.modelo_embedding = texto(
            "NORMALIZADOR_HAB_TEC_EMBEDDING_MODEL", "qwen3-embedding:0.6b"
        )
        catalogo.actualizada_en = _ahora()
        self._persistir(catalogo)
        self._executor.submit(self._vectorizar, catalogo)
        return catalogo.a_dict()

    def obtener(self, identificador: str) -> dict[str, object]:
        return self._obtener(identificador).a_dict()

    def _obtener(self, identificador: str) -> CatalogoHabTec:
        if not re.fullmatch(r"HABTEC_[0-9a-f]{16}", identificador):
            raise KeyError(identificador)
        with self._bloqueo:
            catalogo = self._catalogos.get(identificador)
        if catalogo is None:
            catalogo = self._cargar_persistido(identificador)
            if catalogo is not None:
                with self._bloqueo:
                    self._catalogos[identificador] = catalogo
        if catalogo is None:
            raise KeyError(identificador)
        return catalogo

    def _cargar_persistido(self, identificador: str) -> CatalogoHabTec | None:
        """Recupera un artefacto terminado después de reiniciar la API."""

        directorio = self.base_dir / identificador
        manifest = directorio / "manifest.json"
        if not manifest.is_file():
            return None
        try:
            datos = json.loads(manifest.read_text(encoding="utf-8"))
            if not isinstance(datos, dict):
                return None
            resumen_raw = datos.get("resumen")
            resumen: dict[str, object] = resumen_raw if isinstance(resumen_raw, dict) else {}
            resumen_neo4j = resumen.get("neo4j")
            dimension_raw = resumen.get("dimension_embedding")
            registros = self._leer_registros(directorio / "catalogo_normalizado.jsonl")
        except (OSError, ValueError, json.JSONDecodeError):
            return None
        return CatalogoHabTec(
            id_catalogo=identificador,
            archivo=str(datos.get("archivo") or "catalogo_hab_tec.xlsx"),
            directorio=directorio,
            creado_en=str(datos.get("creado_en") or ""),
            estado=str(datos.get("estado") or "error"),
            hoja=str(datos.get("hoja") or ""),
            registros=registros,
            hallazgos=tuple(
                HallazgoCatalogo(
                    codigo=str(item.get("codigo") or "HAB_TEC_HALLAZGO"),
                    severidad=str(item.get("severidad") or "warning"),
                    mensaje=str(item.get("mensaje") or ""),
                    fila=item.get("fila") if isinstance(item.get("fila"), int) else None,
                    detalle=str(item.get("detalle") or ""),
                )
                for item in datos.get("hallazgos", [])
                if isinstance(item, dict)
            ),
            modelo_embedding=str(resumen.get("modelo_embedding") or ""),
            dimension_embedding=(
                int(dimension_raw) if isinstance(dimension_raw, (int, str)) else 0
            ),
            resumen_neo4j=(
                {
                    "habilidades": int(resumen_neo4j.get("habilidades") or 0),
                    "asociaciones": int(resumen_neo4j.get("asociaciones") or 0),
                }
                if isinstance(resumen_neo4j, dict)
                else None
            ),
            error=str(datos.get("error") or ""),
            actualizada_en=str(datos.get("actualizada_en") or ""),
        )

    @staticmethod
    def _leer_registros(ruta: Path) -> tuple[HabilidadTecnica, ...]:
        if not ruta.is_file():
            return ()
        registros: list[HabilidadTecnica] = []
        for linea in ruta.read_text(encoding="utf-8").splitlines():
            datos = json.loads(linea)
            if not isinstance(datos, dict):
                continue
            registros.append(
                HabilidadTecnica(
                    carrera=str(datos.get("carrera") or ""),
                    id_hab_tec=str(datos.get("id") or ""),
                    nombre=str(datos.get("nombre") or ""),
                    descripcion=str(datos.get("descripcion") or ""),
                    fila=int(datos.get("fila") or 0),
                )
            )
        return tuple(registros)

    @staticmethod
    def _escribir_registros(catalogo: CatalogoHabTec) -> None:
        ruta = catalogo.directorio / "catalogo_normalizado.jsonl"
        with ruta.open("w", encoding="utf-8", newline="\n") as salida:
            for registro in catalogo.registros:
                salida.write(
                    json.dumps(registro.a_dict(incluir_texto=True), ensure_ascii=False) + "\n"
                )

    def _persistir(self, catalogo: CatalogoHabTec) -> None:
        manifest = {
            **catalogo.a_dict(),
            "sha256_fuente": _huella(catalogo.directorio / "catalogo_hab_tec.xlsx"),
        }
        (catalogo.directorio / "manifest.json").write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )

    def _vectorizar(self, catalogo: CatalogoHabTec) -> None:
        try:
            dimension = self._crear_indice(catalogo, catalogo.registros)
            resumen_neo4j = self._publicar_en_neo4j(catalogo.registros, catalogo.id_catalogo)
        except Exception as exc:
            catalogo.estado = "error"
            catalogo.error = f"No se completó la vectorización y publicación: {type(exc).__name__}."
        else:
            catalogo.dimension_embedding = dimension
            catalogo.resumen_neo4j = resumen_neo4j
            catalogo.estado = "vectorizado"
        catalogo.actualizada_en = _ahora()
        self._persistir(catalogo)

    def _crear_indice(
        self, catalogo: CatalogoHabTec, registros: tuple[HabilidadTecnica, ...]
    ) -> int:
        endpoint = texto(
            "NORMALIZADOR_HAB_TEC_EMBEDDING_ENDPOINT", "http://127.0.0.1:11434/api/embed"
        )
        lote = max(1, min(entero("NORMALIZADOR_HAB_TEC_EMBEDDING_BATCH_SIZE", 24), 64))
        indice = catalogo.directorio / "indice"
        indice.mkdir(exist_ok=True)
        metadatos = indice / "metadatos.jsonl"
        vectores = indice / "vectores.f32"
        dimension = 0
        with (
            metadatos.open("w", encoding="utf-8", newline="\n") as archivo_meta,
            vectores.open("wb") as archivo_vectores,
        ):
            for inicio in range(0, len(registros), lote):
                bloque = registros[inicio : inicio + lote]
                embeddings = self._embeddings(
                    endpoint, catalogo.modelo_embedding, [item.texto_indexado for item in bloque]
                )
                if len(embeddings) != len(bloque):
                    raise ErrorCatalogoHabTec(
                        "Ollama devolvió una cantidad inesperada de embeddings."
                    )
                for desplazamiento, (registro, vector) in enumerate(
                    zip(bloque, embeddings, strict=True)
                ):
                    if not vector or not all(math.isfinite(valor) for valor in vector):
                        raise ErrorCatalogoHabTec("Ollama devolvió un embedding inválido.")
                    if dimension == 0:
                        dimension = len(vector)
                    if len(vector) != dimension:
                        raise ErrorCatalogoHabTec(
                            "Ollama devolvió embeddings con dimensiones distintas."
                        )
                    archivo_vectores.write(struct.pack(f"<{dimension}f", *vector))
                    archivo_meta.write(
                        json.dumps(
                            {
                                **registro.a_dict(incluir_texto=True),
                                "indice": inicio + desplazamiento,
                            },
                            ensure_ascii=False,
                        )
                        + "\n"
                    )
        (indice / "manifest.json").write_text(
            json.dumps(
                {
                    "modelo_embedding": catalogo.modelo_embedding,
                    "dimension": dimension,
                    "registros": len(registros),
                    "sha256_catalogo": _huella(catalogo.directorio / "catalogo_normalizado.jsonl"),
                    "formato_vectores": "float32-little-endian",
                },
                ensure_ascii=False,
                indent=2,
            )
            + "\n",
            encoding="utf-8",
        )
        return dimension

    @staticmethod
    def _embeddings(endpoint: str, modelo: str, textos: Iterable[str]) -> list[list[float]]:
        solicitud = urllib.request.Request(
            endpoint,
            data=json.dumps({"model": modelo, "input": list(textos)}).encode("utf-8"),
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        try:
            with urllib.request.urlopen(solicitud, timeout=120) as respuesta:
                datos = json.loads(respuesta.read().decode("utf-8"))
        except (OSError, urllib.error.URLError, json.JSONDecodeError) as exc:
            raise ErrorCatalogoHabTec("No se pudo conectar con Ollama para vectorizar.") from exc
        vectores = datos.get("embeddings") if isinstance(datos, dict) else None
        if not isinstance(vectores, list):
            raise ErrorCatalogoHabTec("Ollama no devolvió embeddings.")
        resultado: list[list[float]] = []
        for vector in vectores:
            if not isinstance(vector, list):
                raise ErrorCatalogoHabTec("Ollama devolvió un embedding inválido.")
            try:
                resultado.append([float(valor) for valor in vector])
            except (TypeError, ValueError) as exc:
                raise ErrorCatalogoHabTec("Ollama devolvió un embedding inválido.") from exc
        return resultado


gestor_catalogos_hab_tec = GestorCatalogosHabTec()
