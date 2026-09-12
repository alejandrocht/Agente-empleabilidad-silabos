"""Extracción estructural DOCX y primitivas curriculares compartidas."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document

from agente.normalizador.empleabilidad.entrada import normalizar_etiqueta
from agente.normalizador.identidad import hashed
from agente.normalizador.modelos import Hallazgo

_PATRON_REFERENCIA_CURRICULAR = r"(?<![A-Z0-9])(?:L\d+|[GE]\d+|[GE]{2,4})(?![A-Z0-9])"
_SECCIONES_HERRAMIENTAS = (
    "recursos de aprendizaje",
    "recursos tecnologicos",
    "recursos tecnológicos",
    "software",
    "herramientas digitales",
)


def _normalizar_modalidad(valor: object) -> str:
    """Reduce la modalidad declarada al único vocabulario permitido por curso.csv."""

    clave = normalizar_etiqueta(valor).replace("_", " ").upper()
    if "PRESENCIAL" in clave:
        return "Presencial"
    if "HIBRID" in clave:
        return "Híbrido"
    if "VIRTUAL" in clave:
        return "Virtual"
    return ""


def _normalizar_naturaleza(valor: object) -> str:
    """Publica la naturaleza académica declarada (Obligatorio/Electivo) en curso.csv."""

    clave = normalizar_etiqueta(valor).replace("_", " ").upper()
    if "OBLIGAT" in clave:
        return "Obligatorio"
    if "ELECTIV" in clave:
        return "Electivo"
    if "ESPECIAL" in clave:
        return "Especialidad"
    if "GENERAL" in clave:
        return "General"
    texto = _texto(valor)
    return texto[:1].upper() + texto[1:].lower() if texto else ""


def _texto(valor: object) -> str:
    return re.sub(r"\s+", " ", str(valor or "")).strip()


def _sin_referencias_curriculares(valor: object) -> str:
    """Remove display-only outcome and competency labels from exported text."""

    return _texto(re.sub(_PATRON_REFERENCIA_CURRICULAR, " ", str(valor or "")))


def _clave(valor: object) -> str:
    return normalizar_etiqueta(valor).replace(" ", "_")


# `_hash_id` stays bound so downstream importers can still bind it from this
# module; it is now only a name for the shared helper.
_hash_id = hashed


def _ids_curriculares(
    carrera: str,
    periodo: str,
    nombre: str,
    codigo_curso: str,
) -> tuple[str, str]:
    """Construye IDs compatibles con el identificador padre de los CSV."""

    codigo = _texto(codigo_curso)
    if codigo:
        return hashed("SIL", codigo), hashed("CUR", codigo)
    return (
        hashed("SIL", carrera, periodo, nombre),
        hashed("CUR", carrera, periodo, nombre),
    )


def _hallazgo(
    codigo: str,
    severidad: str,
    mensaje: str,
    archivo: str,
    detalle: str | None = None,
) -> Hallazgo:
    return Hallazgo(
        codigo=codigo,
        severidad=severidad,  # type: ignore[arg-type]
        mensaje=mensaje,
        hoja=archivo,
        detalle=detalle,
    )


def _filas_tabla(tabla: Any) -> list[list[str]]:
    """Devuelve filas lógicas sin repetir proyecciones de celdas combinadas.

    ``python-docx`` proyecta el texto de la celda que inicia una combinación
    vertical sobre cada fila de continuación. Cuando toda la fila continúa una
    combinación, no representa un registro curricular adicional. También proyecta una celda
    horizontalmente combinada una vez por cada columna de la cuadrícula; esas repeticiones
    consecutivas deben representar una sola celda lógica.
    """

    filas: list[list[str]] = []
    for fila in tabla.rows:
        celdas_fisicas = fila._tr.tc_lst
        if celdas_fisicas and all(_es_continuacion_vertical(celda) for celda in celdas_fisicas):
            continue
        celdas_logicas: list[str] = []
        celda_anterior = None
        for celda in fila.cells:
            if celda._tc is celda_anterior:
                continue
            celdas_logicas.append(_texto(celda.text))
            celda_anterior = celda._tc
        filas.append(celdas_logicas)
    return filas


def _es_continuacion_vertical(celda: Any) -> bool:
    combinacion = celda.tcPr.vMerge
    return combinacion is not None and combinacion.val != "restart"


def _codigos(texto: str, aceptar_no_numericos: bool = True) -> list[str]:
    """Extrae códigos declarados, incluidos formatos curriculares no estándar."""

    valor = _texto(texto).upper()
    numericos = re.findall(r"(?<![A-Z0-9])[GE]\d+(?![A-Z0-9])", valor)
    if not aceptar_no_numericos:
        return list(dict.fromkeys(numericos))
    if numericos:
        palabras_cortas = re.findall(
            r"(?<![A-Z0-9])[A-Z]{2,4}(?![A-Z0-9])",
            valor,
        )
        if len(valor.split()) <= 4:
            return list(dict.fromkeys(numericos + palabras_cortas))
        return list(dict.fromkeys(numericos))

    # Algunas carreras usan referencias como ``EE``. Solo se aceptan códigos
    # alfabéticos cuando la celda es corta; así no se convierten palabras de una
    # descripción extensa en competencias.
    palabras = re.findall(r"(?<![A-Z0-9])[A-Z]{2,4}(?![A-Z0-9])", valor)
    if len(valor.split()) <= 4:
        return list(dict.fromkeys(palabras))
    return []


def _seccion_herramientas(encabezado: str) -> str:
    normalizado = encabezado.replace("_", " ")
    return next(
        (seccion for seccion in _SECCIONES_HERRAMIENTAS if seccion == normalizado),
        "",
    )


def _herramientas_desde_tabla(filas: list[list[str]]) -> list[dict[str, str]]:
    """Toma solo la fila que sigue a un encabezado estructurado exacto."""

    resultado: list[dict[str, str]] = []
    for indice, fila in enumerate(filas[:-1]):
        seccion = _seccion_herramientas(_clave(" ".join(fila)))
        if not seccion:
            continue
        evidencia = _texto(" ".join(filas[indice + 1]))
        if evidencia:
            resultado.append({"seccion": seccion, "texto": evidencia})
    return resultado


def _herramientas_desde_parrafos(doc: Any) -> list[dict[str, str]]:
    """Extrae solo el contenido inmediatamente posterior a un encabezado confiable."""

    resultado: list[dict[str, str]] = []
    parrafos = [_texto(parrafo.text) for parrafo in doc.paragraphs]
    for indice, parrafo in enumerate(parrafos):
        seccion = _seccion_herramientas(_clave(parrafo))
        if not seccion:
            continue
        for candidato in parrafos[indice + 1 : indice + 3]:
            if not candidato or _seccion_herramientas(_clave(candidato)):
                break
            if re.match(r"^(?:[IVXLC]+|\d+)\s*[.)]", candidato, flags=re.IGNORECASE):
                break
            resultado.append({"seccion": seccion, "texto": candidato})
            break
    return resultado


def _deduplicar_evidencias_herramientas(
    evidencias: list[dict[str, str]],
) -> list[dict[str, str]]:
    resultado: list[dict[str, str]] = []
    vistos: set[tuple[str, str]] = set()
    for evidencia in evidencias:
        seccion = _texto(evidencia.get("seccion"))
        texto = _texto(evidencia.get("texto"))
        clave = (seccion.lower(), texto.lower())
        if seccion and texto and clave not in vistos:
            vistos.add(clave)
            resultado.append({"seccion": seccion, "texto": texto})
    return resultado


def _primer_metadata(metadata: dict[str, str], claves: tuple[str, ...]) -> str:
    for clave in claves:
        if metadata.get(clave):
            return metadata[clave]
    return ""


def _unir_metadata(*valores: str) -> str:
    """Conserva valores declarados en un orden estable dentro de una celda CSV."""

    resultado: list[str] = []
    vistos: set[str] = set()
    for valor in valores:
        texto = _texto(valor)
        clave = texto.casefold()
        if texto and clave not in vistos:
            vistos.add(clave)
            resultado.append(texto)
    return " | ".join(resultado)


def _nombre_desde_archivo(nombre: str) -> str:
    base = Path(nombre).stem.replace("_", " ").replace("-", " ")
    return re.sub(r"\s+", " ", base).strip().title()


def _ciclo_desde_ruta(nombre: str) -> str:
    coincidencia = re.search(r"ciclo[_ -]?(\d+)", nombre, re.IGNORECASE)
    return coincidencia.group(1) if coincidencia else ""


def _extraer_docx(
    ruta: Path,
    nombre: str,
    carrera: str,
    periodo: str,
) -> dict[str, object]:
    doc = Document(str(ruta))
    metadata: dict[str, str] = {}
    ultima_etiqueta = ""
    if doc.tables:
        for fila in doc.tables[0].rows:
            valores = [_texto(celda.text) for celda in fila.cells]
            if len(valores) >= 2 and valores[0]:
                ultima_etiqueta = _clave(valores[0])
                metadata[ultima_etiqueta] = _unir_metadata(
                    metadata.get(ultima_etiqueta, ""), valores[1]
                )
            elif len(valores) >= 2 and ultima_etiqueta and valores[1]:
                # Algunos formatos extienden el coordinador en una fila sin
                # etiqueta. Se conserva en orden en vez de descartar nombres.
                metadata[ultima_etiqueta] = _unir_metadata(
                    metadata.get(ultima_etiqueta, ""), valores[1]
                )

    sumilla = ""
    if len(doc.tables) > 1:
        candidato = _texto(
            " ".join(celda.text for fila in doc.tables[1].rows for celda in fila.cells)
        )
        if candidato and _clave(candidato) != "sumilla":
            sumilla = candidato
    logro_general = ""
    logros: list[dict[str, object]] = []
    competencias: list[dict[str, str]] = []
    programa: list[str] = []
    herramientas_evidencia: list[dict[str, str]] = []

    for tabla in doc.tables:
        filas = _filas_tabla(tabla)
        if not filas:
            continue
        encabezado = _clave(" ".join(filas[0]))
        if "competencias_genericas" in encabezado or "competencias_especificas" in encabezado:
            for valores in filas[1:]:
                if len(valores) < 3:
                    continue
                codigos = _codigos(valores[-1])
                if codigos and valores[0] and not re.fullmatch(r"L\d+", valores[0], re.I):
                    competencias.append(
                        {
                            "orden": str(len(competencias) + 1),
                            "nombre": _sin_referencias_curriculares(valores[0]).strip(" ."),
                            "descripcion": _sin_referencias_curriculares(valores[1]),
                            # El código declarado viaja al catálogo final como
                            # `codigo_competencia`.
                            "codigo": codigos[0],
                            "texto_evidencia": _sin_referencias_curriculares(
                                " | ".join(valores[:2])
                            ),
                        }
                    )
        if "logro_de_aprendizaje_general" in encabezado:
            if len(filas) > 1 and filas[1]:
                logro_general = filas[1][0]
            for valores in filas[1:]:
                if len(valores) >= 3 and re.fullmatch(r"L\d+", valores[0], re.I):
                    logros.append(
                        {
                            "orden": str(len(logros) + 1),
                            "descripcion": _sin_referencias_curriculares(valores[1]),
                            "texto_evidencia": _sin_referencias_curriculares(valores[1]),
                        }
                    )
        if "semana" in encabezado and ("tema" in encabezado or "contenido" in encabezado):
            for valores in filas[1:]:
                if valores and re.fullmatch(r"(?:[1-9]|1[0-5])", valores[0]):
                    programa.append(" | ".join(valor for valor in valores[1:3] if valor))
        herramientas_evidencia.extend(_herramientas_desde_tabla(filas))

    herramientas_evidencia.extend(_herramientas_desde_parrafos(doc))

    if not sumilla:
        for tabla in doc.tables:
            filas = _filas_tabla(tabla)
            for indice, valores in enumerate(filas):
                if valores and _clave(valores[0]) == "sumilla" and indice + 1 < len(filas):
                    sumilla = filas[indice + 1][0]
                    break
            if sumilla:
                break

    curso = _primer_metadata(metadata, ("nombre_del_curso", "nombre_curso", "curso", "asignatura"))
    curso = curso or _nombre_desde_archivo(nombre)
    nivel = _primer_metadata(metadata, ("nivel", "ciclo"))
    ciclo = _ciclo_desde_ruta(nombre) or nivel
    texto_relevante = " ".join(
        parte
        for parte in (
            sumilla,
            logro_general,
            " ".join(str(logro["descripcion"]) for logro in logros),
            " ".join(programa),
        )
        if parte
    )
    texto_fuente = _sin_referencias_curriculares(
        " ".join(
            [parrafo.text for parrafo in doc.paragraphs]
            + [celda.text for tabla in doc.tables for fila in tabla.rows for celda in fila.cells]
        )
    )
    id_silabo, id_curso = _ids_curriculares(
        carrera,
        periodo,
        nombre,
        _primer_metadata(metadata, ("codigo_del_curso", "codigo", "cod_asignatura")),
    )
    return {
        "id_silabo": id_silabo,
        "id_curso": id_curso,
        "carrera": carrera,
        "periodo": periodo,
        "origen": {"archivo": nombre, "formato": "docx"},
        "datos": {
            "curso": curso,
            "ciclo": ciclo,
            "nombre_curso": curso,
            "coordinador": _primer_metadata(metadata, ("coordinador", "coordinador_del_curso")),
            "creditos": _primer_metadata(metadata, ("creditos", "creditos_academicos")),
            "nivel": nivel,
            "tipo_curso": _normalizar_naturaleza(
                _primer_metadata(
                    metadata,
                    ("tipo_de_asignatura", "tipo_asignatura", "naturaleza"),
                )
            ),
            "modalidad": _normalizar_modalidad(
                _primer_metadata(
                    metadata,
                    ("modalidad", "modalidad_de_estudios", "modalidad_de_ensenanza"),
                )
            ),
            "codigo_curso": _primer_metadata(
                metadata,
                ("codigo_del_curso", "codigo", "cod_asignatura"),
            ),
            "sumilla": sumilla,
            "logro_general": logro_general,
            "logros_especificos": logros,
            "competencias_declaradas": competencias,
            "programa_analitico": programa,
            "herramientas_evidencia": _deduplicar_evidencias_herramientas(herramientas_evidencia),
            "texto_relevante": texto_relevante,
            "texto_fuente": texto_fuente,
        },
    }
