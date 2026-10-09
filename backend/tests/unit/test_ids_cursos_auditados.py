"""Disambiguación de los cuatro códigos repetidos auditados en 2026-2."""

from __future__ import annotations

import csv
from pathlib import Path

import pytest
from docx import Document

from agente.normalizador.silabos import extraccion_pdf
from agente.normalizador.silabos.extraccion_curricular import (
    _extraer_docx,
    _hash_id,
    _ids_curriculares,
)
from agente.normalizador.silabos.salida_catalogos import construir_catalogos_curriculares

CASOS = (
    ("ARQUITECTURA", "700087", "DESARROLLO DE PROYECTO", "GESTION DE PROYECTOS II"),
    ("ECONOMIA", "530004", "METODOS NO PARAMETRICOS", "ESTADISTICA APLICADA II"),
    ("INGENIERIA_CIVIL", "710059", "FUNDAMENTOS DEL PLANEAMIENTO URBANO Y REGIONAL", "HIDRAULICA"),
    (
        "DERECHO",
        "7410",
        "DERECHO CIVIL I PRINCIPIOS GENERALES Y PERSONAS NATURALES",
        "DERECHO PROCESAL CIVIL I",
    ),
)


@pytest.mark.parametrize("carrera,codigo,curso,contraparte", CASOS)
def test_separa_curso_y_silabo_sin_cambiar_ids_de_contraparte(
    carrera: str,
    codigo: str,
    curso: str,
    contraparte: str,
) -> None:
    ids = _ids_curriculares(carrera, "2026-2", f"Ciclo_08/{curso.replace(' ', '_')}.docx", codigo)
    originales = (_hash_id("SIL", codigo, "2026-2"), _hash_id("CUR", codigo))
    assert ids[0] != originales[0]
    assert ids[1] != originales[1]
    assert _ids_curriculares(carrera, "2026-2", f"{contraparte}.docx", codigo) == originales
    assert _ids_curriculares("OTRA_CARRERA", "2026-2", f"{curso}.docx", codigo) == originales
    assert _ids_curriculares(carrera, "2026-2", f"{curso}.docx", "CODIGO_CORREGIDO") == (
        _hash_id("SIL", "CODIGO_CORREGIDO", "2026-2"),
        _hash_id("CUR", "CODIGO_CORREGIDO"),
    )


@pytest.mark.parametrize("carrera,codigo,curso,contraparte", CASOS)
def test_hash_estable_por_nombre_y_carrera_ignora_ruta_y_formato(
    carrera: str,
    codigo: str,
    curso: str,
    contraparte: str,
) -> None:
    primero = _ids_curriculares(carrera, "2026-2", f"Ciclo_03/{curso}.docx", codigo)
    variante = _ids_curriculares(
        carrera.lower().replace("_", " "),
        "2026-2",
        f"otra/carpeta/{curso.lower().replace(' ', '_')}.pdf",
        codigo,
    )
    assert variante == primero
    renombrado = _ids_curriculares(
        carrera, "2026-2", "archivo_renombrado.pdf", codigo, nombre_curso=curso
    )
    assert renombrado == primero
    siguiente = _ids_curriculares(carrera, "2027-1", f"{curso}.docx", codigo)
    assert siguiente[0] != primero[0]
    assert siguiente[1] == primero[1]


@pytest.mark.parametrize("carrera,codigo,curso,contraparte", CASOS)
def test_docx_y_csv_conservan_codigo_fuente_y_relaciones_distintas(
    tmp_path: Path,
    carrera: str,
    codigo: str,
    curso: str,
    contraparte: str,
) -> None:
    registros = []
    for indice, nombre in enumerate((curso, contraparte)):
        ruta = tmp_path / f"documento_{indice}.docx"
        documento = Document()
        metadata = documento.add_table(rows=2, cols=2)
        metadata.cell(0, 0).text = "Asignatura"
        metadata.cell(0, 1).text = nombre
        metadata.cell(1, 0).text = "Código"
        metadata.cell(1, 1).text = codigo
        documento.save(ruta)
        registro = _extraer_docx(ruta, ruta.name, carrera, "2026-2")
        # Añadir evidencia mínima para verificar las referencias derivadas de los IDs.
        datos = registro["datos"]
        assert isinstance(datos, dict)
        assert datos["codigo_curso"] == codigo
        datos["competencias_declaradas"] = [
            {
                "nombre": "Competencia fuente",
                "descripcion": "Descripción fuente.",
                "codigo": "E1",
                "tipo": "especifica",
            },
        ]
        datos["logro_general"] = "Un logro literal de la fuente."
        registros.append(registro)

    assert len({r["id_curso"] for r in registros}) == 2
    assert len({r["id_silabo"] for r in registros}) == 2
    assert registros[1]["id_curso"] == _hash_id("CUR", codigo)
    salida = tmp_path / "salidas"
    construir_catalogos_curriculares(registros, salida, carrera=carrera, periodo_academico="2026-2")

    def leer(nombre: str) -> list[dict[str, str]]:
        with (salida / nombre).open(encoding="utf-8-sig", newline="") as archivo:
            return list(csv.DictReader(archivo))

    cursos, silabos, coberturas = (
        leer("curso.csv"),
        leer("silabo.csv"),
        leer("cobertura_curricular.csv"),
    )
    assert len(cursos) == len(silabos) == len(coberturas) == 2
    assert {fila["codigo_curso"] for fila in cursos} == {codigo}
    assert {fila["id_curso"] for fila in silabos} == {r["id_curso"] for r in registros}
    assert {fila["id_silabo"] for fila in coberturas} == {r["id_silabo"] for r in registros}
    assert {fila["id_curso"] for fila in coberturas} == {r["id_curso"] for r in registros}
    assert len({fila["id_logro"] for fila in coberturas}) == 2


@pytest.mark.parametrize("carrera,codigo,curso,contraparte", CASOS)
def test_pdf_usa_nombre_declarado_aunque_se_renombre_archivo(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    carrera: str,
    codigo: str,
    curso: str,
    contraparte: str,
) -> None:
    class Pagina:
        def extract_text(self, **kwargs: object) -> str:
            return f"I. Identificación\nAsignatura: {curso}\nCódigo: {codigo}\nNivel: 8\n"

    class Lector:
        pages = [Pagina()]

    registro = extraccion_pdf._extraer_pdf(
        tmp_path / "renombrado.pdf",
        "renombrado.pdf",
        carrera,
        "2026-2",
        pdf_reader=lambda *_args: Lector(),
        geometria_pdf_pagina=lambda *_args: ([], []),
    )
    assert (registro["id_silabo"], registro["id_curso"]) == _ids_curriculares(
        carrera,
        "2026-2",
        f"{curso}.docx",
        codigo,
    )
    datos = registro["datos"]
    assert isinstance(datos, dict)
    assert datos["codigo_curso"] == codigo


def test_ids_de_los_ocho_cursos_son_distintos() -> None:
    identidades = [
        _ids_curriculares(carrera, "2026-2", f"{nombre}.docx", codigo)
        for carrera, codigo, curso, contraparte in CASOS
        for nombre in (curso, contraparte)
    ]
    assert len({silabo for silabo, _curso in identidades}) == 8
    assert len({curso for _silabo, curso in identidades}) == 8
