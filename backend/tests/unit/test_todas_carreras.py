"""Una ejecución multicarrera: procedencia explícita, IDs y cobertura completa."""

from __future__ import annotations

import csv
import json
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from agente.normalizador.excepciones import CancelacionSolicitada
from agente.normalizador.silabos.cactus_archivos import CactusExtractorError
from agente.normalizador.silabos.cactus_navegacion import CactusAuthenticationError
from agente.normalizador.silabos.entrada import (
    CARRERAS_ULIMA,
    carrera_del_archivo,
    normalizar_carrera,
    validar_archivo,
)
from agente.normalizador.silabos.extraccion_curricular import distinguir_cursos_multicarrera
from agente.normalizador.silabos.fuente_cactus import CactusExtractor, ResultadoExtraccionCactus
from agente.normalizador.silabos.limpieza import limpiar_archivo


def _zip(ruta: Path, nombres: list[str], contenido: bytes = b"documento") -> Path:
    with ZipFile(ruta, "w") as paquete:
        for nombre in nombres:
            paquete.writestr(nombre, contenido)
    return ruta


@pytest.mark.parametrize("carrera", CARRERAS_ULIMA)
def test_reconoce_carpetas_de_las_catorce_carreras(carrera: str) -> None:
    assert carrera_del_archivo(f"silabos/{carrera}/2026-2/Ciclo_01/a.docx", "2026-2") == (
        normalizar_carrera(carrera)
    )


def test_acepta_alias_sistemas_y_carpetas_sin_periodo() -> None:
    assert carrera_del_archivo("INGENIERIA_SISTEMAS/Ciclo_01/a.pdf", "2026-2") == (
        "INGENIERIA_DE_SISTEMAS"
    )


@pytest.mark.parametrize(
    "nombre",
    [
        "Ciclo_01/a.docx",
        "Marketing/2026-1/a.pdf",
        "Marketing/Arquitectura/a.docx",
        "../Marketing/a.docx",
    ],
)
def test_rechaza_carpetas_ambiguas_periodos_mezclados_y_rutas_inseguras(
    tmp_path: Path, nombre: str
) -> None:
    resultado = validar_archivo(_zip(tmp_path / "entrada.zip", [nombre]), "TODAS", "2026-2")
    assert not resultado.valida
    assert any(h.severidad == "error" for h in resultado.hallazgos)


def test_todas_requiere_zip(tmp_path: Path) -> None:
    ruta = tmp_path / "a.docx"
    ruta.write_bytes(b"docx")
    resultado = validar_archivo(ruta, "Todas las carreras", "2026-2")
    assert not resultado.valida
    assert "TODAS_CARRERAS_REQUIERE_ZIP" in {h.codigo for h in resultado.hallazgos}


def test_admite_mas_de_quinientos_silabos_y_mismo_nombre_en_dos_carreras(tmp_path: Path) -> None:
    nombres = [
        f"{carrera}/2026-2/Ciclo_01/curso_{i}.docx"
        for carrera in ("Marketing", "Arquitectura")
        for i in range(251)
    ]
    ruta = _zip(tmp_path / "entrada.zip", nombres)
    resultado = validar_archivo(ruta, "TODAS", "2026-2")
    assert resultado.valida
    assert len(resultado.archivos) == 502
    individual = validar_archivo(ruta, "Marketing", "2026-2")
    assert not individual.valida
    assert "LIMITE_ARCHIVOS_EXCEDIDO" in {h.codigo for h in individual.hallazgos}


def test_cursos_compartidos_tienen_ids_estables_y_los_unicos_conservan_id() -> None:
    registros = [
        {"id_curso": "CUR_1", "id_silabo": "SIL_1", "carrera": "MARKETING"},
        {"id_curso": "CUR_1", "id_silabo": "SIL_1", "carrera": "ARQUITECTURA"},
        {"id_curso": "CUR_2", "id_silabo": "SIL_2", "carrera": "MARKETING"},
    ]
    repeticion = deepcopy(registros)
    distinguir_cursos_multicarrera(registros)
    distinguir_cursos_multicarrera(repeticion)
    assert registros == repeticion
    assert len({r["id_curso"] for r in registros}) == 3
    assert len({r["id_silabo"] for r in registros}) == 3
    assert registros[-1]["id_curso"] == "CUR_2"
    assert registros[-1]["id_silabo"] == "SIL_2"
    distinguir_cursos_multicarrera(registros)
    assert registros == repeticion


def test_zip_multicarrera_materializa_un_paquete_con_referencias_coherentes(tmp_path: Path) -> None:
    fuente = tmp_path / "fuente.docx"
    documento = Document()
    metadata = documento.add_table(rows=2, cols=2)
    metadata.cell(0, 0).text, metadata.cell(0, 1).text = "Curso", "Estadística"
    metadata.cell(1, 0).text, metadata.cell(1, 1).text = "Código", "10001"
    competencia = documento.add_table(rows=2, cols=3)
    for i, texto in enumerate(("Competencias genéricas", "Descripción", "Código")):
        competencia.cell(0, i).text = texto
    for i, texto in enumerate(("Análisis estadístico", "Analiza datos.", "G1")):
        competencia.cell(1, i).text = texto
    logro = documento.add_table(rows=2, cols=3)
    for i, texto in enumerate(("Logro de aprendizaje general", "Descripción", "Competencias")):
        logro.cell(0, i).text = texto
    for i, texto in enumerate(("L1", "Analiza distribuciones estadísticas.", "G1")):
        logro.cell(1, i).text = texto
    documento.save(fuente)
    ruta = _zip(
        tmp_path / "entrada.zip",
        [f"{c}/2026-2/Ciclo_01/ESTADISTICA.docx" for c in ("Marketing", "Arquitectura")],
        fuente.read_bytes(),
    )
    validacion = validar_archivo(ruta, "TODAS", "2026-2")
    assert validacion.valida
    destino = tmp_path / "ejecucion"
    resultado = limpiar_archivo(ruta, destino, validacion)
    assert resultado.registros == 2
    assert resultado.publicable
    staging = [
        json.loads(linea) for linea in (destino / "limpios/silabos.jsonl").read_text().splitlines()
    ]
    assert {r["carrera"] for r in staging} == {"MARKETING", "ARQUITECTURA"}
    filas = {}
    for nombre in ("curso", "silabo", "cobertura_curricular"):
        with (destino / f"salidas/{nombre}.csv").open(encoding="utf-8-sig") as archivo:
            filas[nombre] = list(csv.DictReader(archivo))
    assert len({r["id_curso"] for r in filas["curso"]}) == 2
    assert len({r["id_carrera"] for r in filas["curso"]}) == 2
    assert len({r["id_silabo"] for r in filas["silabo"]}) == 2
    assert {r["id_curso"] for r in filas["cobertura_curricular"]} == {
        r["id_curso"] for r in filas["curso"]
    }
    assert {r["id_silabo"] for r in filas["cobertura_curricular"]} == {
        r["id_silabo"] for r in filas["silabo"]
    }


def _resultado(carrera: str, directorio: Path) -> ResultadoExtraccionCactus:
    archivo = directorio / "Ciclo_01/curso.docx"
    archivo.parent.mkdir(parents=True)
    archivo.write_bytes(b"docx")
    return ResultadoExtraccionCactus(
        carrera=carrera,
        periodo="2026-2",
        cursos_encontrados=1,
        archivos_descargados=1,
        archivos_procesables=1,
        sin_silabo=0,
        fetch_fallidos=0,
        sesiones_fallidas=0,
        archivos_no_soportados=0,
        archivos=(archivo,),
        errores=(),
    )


def _extraer_todas(extractor: CactusExtractor, tmp_path: Path, **kwargs):
    return extractor._extraer_todas(
        periodo="2026-2",
        usuario="usuario",
        contrasena="secreto",
        directorio_salida=tmp_path / "descargas",
        directorio_perfil=tmp_path / "perfil",
        al_actualizar_progreso=kwargs.get("progreso"),
        cancelada=kwargs.get("cancelada"),
    )


def test_cactus_recorre_todas_y_acumula_progreso_y_carpetas(monkeypatch, tmp_path: Path) -> None:
    extractor = CactusExtractor()
    llamadas = []
    eventos = []

    def extraer(**kwargs):
        llamadas.append(kwargs["carrera"])
        kwargs["al_actualizar_progreso"]({"cursos_encontrados": 1, "archivos_descargados": 1})
        return _resultado(kwargs["carrera"], kwargs["directorio_salida"])

    monkeypatch.setattr(extractor, "extraer", extraer)
    resultado = _extraer_todas(extractor, tmp_path, progreso=eventos.append)
    assert llamadas == list(CARRERAS_ULIMA)
    assert resultado.completa
    assert resultado.cursos_encontrados == resultado.archivos_descargados == 14
    assert eventos[-1]["carreras_procesadas"] == 14
    assert eventos[-2]["archivos_descargados"] == 14
    for archivo, carrera in zip(resultado.archivos, CARRERAS_ULIMA, strict=True):
        assert carrera_del_archivo(archivo.as_posix(), "2026-2") == normalizar_carrera(carrera)


def test_cactus_carrera_fallida_bloquea_publicacion_aunque_otros_conteos_cuadren(
    monkeypatch, tmp_path: Path
) -> None:
    extractor = CactusExtractor()

    def extraer(**kwargs):
        if kwargs["carrera"] == "Arquitectura":
            raise CactusExtractorError("FUENTE_FALLIDA", "No se pudo descargar la carrera")
        return _resultado(kwargs["carrera"], kwargs["directorio_salida"])

    monkeypatch.setattr(extractor, "extraer", extraer)
    resultado = _extraer_todas(extractor, tmp_path)
    assert resultado.cursos_encontrados == resultado.archivos_descargados == 13
    assert not resultado.completa
    assert resultado.carreras_incompletas == ("Arquitectura",)
    assert resultado.errores[0]["carrera"] == "Arquitectura"
    assert resultado.a_dict(tmp_path)["completa"] is False


@pytest.mark.parametrize("autenticacion", [True, False])
def test_cactus_detiene_todas_al_fallar_autenticacion_o_cancelarse(
    monkeypatch, tmp_path: Path, autenticacion: bool
) -> None:
    extractor = CactusExtractor()
    llamadas = []

    def extraer(**kwargs):
        llamadas.append(kwargs["carrera"])
        if autenticacion:
            raise CactusAuthenticationError("Sesión inválida")
        return _resultado(kwargs["carrera"], kwargs["directorio_salida"])

    monkeypatch.setattr(extractor, "extraer", extraer)
    error = CactusAuthenticationError if autenticacion else CancelacionSolicitada
    with pytest.raises(error):
        _extraer_todas(extractor, tmp_path, cancelada=lambda: bool(llamadas) and not autenticacion)
    assert llamadas == ["Administración"]
