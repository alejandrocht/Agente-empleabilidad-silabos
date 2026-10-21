"""Closed HAB_TEC vocabulary at retrieval, inference and publication boundaries."""

from __future__ import annotations

import csv
import json
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pydantic import ValidationError

from agente.config.settings import configuracion_normalizador_curricular
from agente.normalizador.catalogo_hab_tec import GestorCatalogosHabTec
from agente.normalizador.empleabilidad.hab_tec_retriever import (
    ErrorHabTecRetriever,
    HabTecRetriever,
    cargar_retriever_hab_tec,
)
from agente.normalizador.excepciones import CancelacionSolicitada
from agente.normalizador.silabos import analista_tecnico, salida_catalogos

EMBEDDING_REAL = HabTecRetriever._embedding

CARRERA = "Ingeniería de Sistemas"
NOMBRE = "Arquitectura de software"
DESCRIPCION = "Diseña estructuras y patrones técnicos."
LOGRO = "Aplica Arquitectura de software y Analítica comercial para resolver problemas técnicos."


def _vectorizar(
    base: Path, monkeypatch: pytest.MonkeyPatch, id_habilidad: str = "HAB_TEC_001"
) -> Path:
    fuente = base.parent / "entrada.xlsx"
    libro = Workbook()
    hoja = libro.active
    assert hoja is not None
    hoja.append(["Carrera", "id", "nombre", "descripcion"])
    hoja.append([CARRERA, id_habilidad, NOMBRE, DESCRIPCION])
    hoja.append([CARRERA, "HAB_TEC_002", "Redes", "Configura redes informáticas."])
    hoja.append(["Marketing", id_habilidad, "Analítica comercial", "Analiza campañas."])
    libro.save(fuente)
    gestor = GestorCatalogosHabTec(base, publicar_en_neo4j=lambda *_args: {})
    estado = gestor.crear_desde_archivo(fuente.name, fuente)
    catalogo = gestor._obtener(str(estado["id_catalogo"]))
    catalogo.modelo_embedding = "modelo-del-indice"
    monkeypatch.setattr(gestor, "_embeddings", lambda *_args: [[1, 0], [0, 1], [1, 0]])
    gestor._vectorizar(catalogo)
    assert catalogo.estado == "vectorizado"
    return catalogo.directorio


@pytest.fixture
def indice(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    base = tmp_path / "hab_tec"
    monkeypatch.setenv("NORMALIZADOR_HAB_TEC_DIR", str(base))
    monkeypatch.setenv("NORMALIZADOR_HAB_TEC_EMBEDDING_MODEL", "otro-modelo-configurado")
    monkeypatch.setenv("NORMALIZADOR_HAB_TEC_RETRIEVAL_MIN_SIMILARITY", "0.5")
    monkeypatch.setattr(HabTecRetriever, "_embedding", lambda *_args: (1.0, 0.0))
    return _vectorizar(base, monkeypatch)


def _registro(id_silabo: str = "SIL_1") -> dict[str, Any]:
    return {
        "id_curso": "CUR_" + id_silabo,
        "id_silabo": id_silabo,
        "carrera": "INGENIERIA_DE_SISTEMAS",
        "periodo": "2026-2",
        "datos": {
            "nombre_curso": NOMBRE,
            "logro_general": LOGRO,
            "competencias_declaradas": [],
        },
    }


def _propuesta(ref: str | None = "HAB_TEC_001") -> dict[str, Any]:
    return {
        "catalogo_ref": ref,
        "nombre_competencia": "Nombre inventado por el modelo",
        "descripcion_breve_competencia": "Descripción alterada por el modelo.",
        "logros": [LOGRO],
        "evidencia": [{"fuente": "logro", "fragmento": LOGRO}],
        "justificacion": "El logro sustenta la habilidad seleccionada.",
    }


def _modelo(monkeypatch: pytest.MonkeyPatch, respuestas: list[dict[str, Any] | None]) -> list[Any]:
    llamadas: list[Any] = []

    class Modelo:
        def with_structured_output(self, *_args: Any, **_kwargs: Any) -> Modelo:
            return self

        def invoke(self, mensajes: Any) -> dict[str, Any] | None:
            llamadas.append(mensajes)
            return deepcopy(respuestas.pop(0) if len(respuestas) > 1 else respuestas[0])

    monkeypatch.setattr(analista_tecnico, "obtener_llm", lambda *_args, **_kwargs: Modelo())
    return llamadas


def _inferir(registros: list[dict[str, Any]], **kwargs: Any) -> list[dict[str, object]]:
    return analista_tecnico.inferir_competencias_tecnicas(
        registros,
        configuracion_normalizador_curricular(
            {
                "NORMALIZADOR_CURRICULAR_LLM": "true",
                "NORMALIZADOR_CURRICULAR_LLM_TIMEOUT_SECONDS": "120",
                "NORMALIZADOR_CURRICULAR_LLM_MAX_RETRIES": "2",
                "NORMALIZADOR_CURRICULAR_LLM_BATCH_SIZE": "8",
                "NORMALIZADOR_CURRICULAR_LLM_TEMPERATURE": "0",
                "NORMALIZADOR_CURRICULAR_ANALYST_REASONING_EFFORT": "medium",
            }
        ),
        **kwargs,
    )


def test_recupera_por_carrera_y_usa_modelo_del_indice(indice: Path) -> None:
    recuperador = cargar_retriever_hab_tec()
    assert recuperador is not None
    assert recuperador.modelo == "modelo-del-indice"
    assert recuperador.id_catalogo == indice.name
    assert [m.nombre for m in recuperador.buscar(LOGRO, "INGENIERIA_DE_SISTEMAS")] == [NOMBRE]


def test_catalogo_del_pull_windows_con_crlf_conserva_hash_del_indice(indice):
    anterior = cargar_retriever_hab_tec()
    fuente = indice / "catalogo_normalizado.jsonl"
    fuente.write_bytes(fuente.read_bytes().replace(b"\n", b"\r\n"))
    actual = cargar_retriever_hab_tec()
    assert anterior is not None and actual is not None
    assert actual.sha256_catalogo == anterior.sha256_catalogo
    assert actual.documentos == anterior.documentos


def test_crlf_no_oculta_alteraciones_del_catalogo(indice):
    fuente = indice / "catalogo_normalizado.jsonl"
    fuente.write_bytes(
        fuente.read_bytes().replace(b"\n", b"\r\n").replace(b"HAB_TEC_001", b"HAB_TEC_999")
    )
    with pytest.raises(ErrorHabTecRetriever, match="no corresponde"):
        cargar_retriever_hab_tec()


@pytest.mark.parametrize("motivo", ["sin_catalogo", "hash_distinto", "faltan_carreras"])
def test_comprobacion_explica_por_que_rechaza_catalogo(indice, monkeypatch, motivo):
    from api import servidor

    monkeypatch.setenv("NORMALIZADOR_CURRICULAR_LLM", "true")
    if motivo == "sin_catalogo":
        monkeypatch.setenv("NORMALIZADOR_HAB_TEC_DIR", str(indice.parent / "otra_carpeta"))
    elif motivo == "hash_distinto":
        fuente = indice / "catalogo_normalizado.jsonl"
        fuente.write_bytes(fuente.read_bytes().replace(b"HAB_TEC_001", b"HAB_TEC_999"))
    monkeypatch.setattr(
        analista_tecnico, "obtener_llm", lambda *_a, **_k: pytest.fail("No consultar LLM")
    )
    respuesta = TestClient(servidor.app, client=("127.0.0.1", 0)).post(cli_ruta_comprobacion())
    assert respuesta.status_code == 503
    detalle = respuesta.json()["detail"]
    if motivo == "faltan_carreras":
        assert "Administración" in detalle["carreras_faltantes"]
        assert "Marketing" not in detalle["carreras_faltantes"]
    else:
        assert (
            detalle["motivo"]
            == {
                "sin_catalogo": "CATALOGO_NO_ENCONTRADO",
                "hash_distinto": "CATALOGO_HASH_DISTINTO",
            }[motivo]
        )


def test_modelo_selecciona_referencia_y_python_copia_campos_exactos(
    indice: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    llamadas = _modelo(monkeypatch, [{"competencias": [_propuesta()]}])
    propuestas = _inferir([_registro()])
    assert len(propuestas) == 1
    assert propuestas[0]["catalogo_ref"] == "HAB_TEC_001"
    assert propuestas[0]["nombre_competencia"] == NOMBRE
    assert propuestas[0]["descripcion_breve_competencia"] == DESCRIPCION
    assert propuestas[0]["catalogo_id"] == indice.name
    assert propuestas[0]["estado_aprobacion"] == "PENDIENTE_APROBACION"
    payload = json.loads(llamadas[0][1][1].split("\n", 1)[1])
    assert [c["catalogo_ref"] for c in payload["catalog_context"]] == ["HAB_TEC_001"]
    assert payload["syllabus_context"]["logros"] == [{"texto": LOGRO}]


def test_acepta_habilidad_sustentada_sin_nombre_explicito(indice, monkeypatch):
    registro = _registro()
    logro = "Diseña estructuras y patrones mantenibles."
    registro["datos"]["logro_general"] = logro
    propuesta = {
        **_propuesta(),
        "logros": [logro],
        "evidencia": [{"fuente": "logro", "fragmento": logro}],
    }
    llamadas = _modelo(monkeypatch, [{"competencias": [propuesta]}])
    trazas = []
    assert len(_inferir([registro], al_actualizar_progreso_silabo=trazas.append)) == 1
    assert len(llamadas) == 1
    assert trazas[-1].latencia_modelo_ms is not None
    assert trazas[-1].estado_analisis == "completado"


def test_modelo_no_puede_devolver_dos_habilidades_por_silabo(indice, monkeypatch):
    _modelo(monkeypatch, [{"competencias": [_propuesta(), _propuesta()]}])
    with pytest.raises(ValidationError):
        _inferir([_registro()])


def test_registros_repetidos_no_agregan_otra_habilidad_al_mismo_silabo(indice, monkeypatch):
    llamadas = _modelo(monkeypatch, [{"competencias": [_propuesta()]}])
    assert len(_inferir([_registro(), _registro()])) == 1
    assert len(llamadas) == 1


def test_exportacion_rechaza_segunda_habilidad_aun_aprobada(indice, monkeypatch, tmp_path):
    _modelo(monkeypatch, [{"competencias": [_propuesta()]}])
    registro = _registro()
    registro["datos"]["programa_analitico_detalle"] = [{"tema": "Redes", "contenido": ""}]
    primera = {**_inferir([registro])[0], "estado_aprobacion": "APROBADA"}
    segunda = {
        **primera,
        "catalogo_ref": "HAB_TEC_002",
        "nombre_competencia": "Redes",
        "descripcion_breve_competencia": "Configura redes informáticas.",
    }
    with pytest.raises(ValueError, match="una habilidad técnica por sílabo"):
        salida_catalogos.construir_catalogos_curriculares(
            [registro],
            tmp_path,
            carrera="INGENIERIA_DE_SISTEMAS",
            periodo_academico="2026-2",
            competencias_tecnicas=[primera, segunda],
        )
    assert not (tmp_path / "catalogo_habilidades.csv").exists()


def test_exportacion_acepta_habilidad_semantica_aprobada(indice, monkeypatch, tmp_path):
    _modelo(monkeypatch, [{"competencias": [_propuesta()]}])
    registro = _registro()
    aprobada = {**_inferir([registro])[0], "estado_aprobacion": "APROBADA"}
    logro = "Diseña estruturas y patrones mantenibles."
    registro["datos"]["logro_general"] = logro
    aprobada["logros"] = [logro]
    aprobada["evidencia"] = [{"fuente": "logro", "fragmento": logro}]
    salida_catalogos.construir_catalogos_curriculares(
        [registro],
        tmp_path,
        carrera="INGENIERIA_DE_SISTEMAS",
        periodo_academico="2026-2",
        competencias_tecnicas=[aprobada],
    )
    assert (tmp_path / "catalogo_habilidades.csv").exists()


@pytest.mark.parametrize(
    "nombre,texto,esperado",
    [
        ("Analítica comercial", "Aplica ANALITICA   COMERCIAL.", True),
        ("Redes", "Estudia paredes estructurales.", False),
        ("C", "Programa con C++.", False),
        ("C++", "Programa con C++.", True),
    ],
)
def test_match_explicito_no_confunde_palabras_o_lenguajes(nombre, texto, esperado):
    candidato = analista_tecnico.CandidatoTecnico("HAB_TEC_001", CARRERA, nombre, "Descripción", 1)
    assert (
        analista_tecnico.habilidad_mencionada(candidato, {"logros": [{"texto": texto}]}) is esperado
    )


@pytest.mark.parametrize("ref", [None, "", "HAB_TEC_999", "HAB_TEC_002"])
def test_rechaza_inventadas_y_referencias_no_recuperadas(
    indice: Path, monkeypatch: pytest.MonkeyPatch, ref: str | None
) -> None:
    llamadas = _modelo(monkeypatch, [{"competencias": [_propuesta(ref)]}])
    auditoria: list[dict[str, object]] = []
    assert _inferir([_registro()], auditoria=auditoria, exigir_habilidad=False) == []
    assert len(llamadas) == 1
    assert auditoria[0]["codigo"] == "HABILIDAD_FUERA_CATALOGO"


def test_reintento_puede_recuperar_referencia_valida(
    indice: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    llamadas = _modelo(
        monkeypatch,
        [
            {"competencias": [_propuesta(None)]},
            {"competencias": [_propuesta()]},
        ],
    )
    assert len(_inferir([_registro()])) == 1
    assert len(llamadas) == 2


def test_misma_habilidad_conserva_evidencia_y_cobertura_de_cada_silabo(
    indice: Path, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    _modelo(monkeypatch, [{"competencias": [_propuesta()]}])
    registros = [_registro(), _registro("SIL_2")]
    propuestas = _inferir(registros)
    assert [p["id_silabo"] for p in propuestas] == ["SIL_1", "SIL_2"]
    assert propuestas[0]["id_propuesta"] != propuestas[1]["id_propuesta"]
    salida = tmp_path / "salida"
    salida_catalogos.construir_catalogos_curriculares(
        registros,
        salida,
        carrera="INGENIERIA_DE_SISTEMAS",
        periodo_academico="2026-2",
        competencias_tecnicas=[{**p, "estado_aprobacion": "APROBADA"} for p in propuestas],
    )
    with (salida / "catalogo_habilidades.csv").open(encoding="utf-8-sig") as archivo:
        habilidades = list(csv.DictReader(archivo))
    assert len(habilidades) == 1
    assert habilidades[0]["id_habilidad"] == "HAB_TEC_001"
    assert habilidades[0]["nombre_habilidad"] == NOMBRE
    assert habilidades[0]["desc_breve"] == DESCRIPCION
    with (salida / "cobertura_curricular.csv").open(encoding="utf-8-sig") as archivo:
        assert {f["id_silabo"] for f in csv.DictReader(archivo)} == {"SIL_1", "SIL_2"}


def test_todas_selecciona_por_carrera_y_exporta_variantes_exactas_del_mismo_id(
    indice: Path, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    llamadas = _modelo(monkeypatch, [{"competencias": [_propuesta()]}])
    marketing = _registro("SIL_MARKETING")
    marketing["carrera"] = "MARKETING"
    registros = [_registro(), marketing]
    propuestas = _inferir(registros)
    assert len(llamadas) == len(propuestas) == 2
    assert [p["nombre_competencia"] for p in propuestas] == [NOMBRE, "Analítica comercial"]
    for llamada, nombre in zip(llamadas, [NOMBRE, "Analítica comercial"], strict=True):
        payload = json.loads(llamada[1][1].split("\n", 1)[1])
        assert len(payload["catalog_context"]) == 1
        assert payload["catalog_context"][0]["nombre"] == nombre
    salida = tmp_path / "salida"
    aprobadas = [{**p, "estado_aprobacion": "APROBADA"} for p in propuestas]
    resumen = salida_catalogos.construir_catalogos_curriculares(
        registros,
        salida,
        carrera="TODAS",
        periodo_academico="2026-2",
        competencias_tecnicas=aprobadas,
    )
    assert resumen["carreras"] == ["INGENIERIA_DE_SISTEMAS", "MARKETING"]
    with (salida / "catalogo_habilidades.csv").open(encoding="utf-8-sig") as archivo:
        habilidades = list(csv.DictReader(archivo))
    assert len(habilidades) == 2
    assert {h["id_habilidad"] for h in habilidades} == {"HAB_TEC_001"}
    assert len({h["id_carrera"] for h in habilidades}) == 2
    assert {(h["nombre_habilidad"], h["desc_breve"]) for h in habilidades} == {
        (NOMBRE, DESCRIPCION),
        ("Analítica comercial", "Analiza campañas."),
    }
    with (salida / "cobertura_curricular.csv").open(encoding="utf-8-sig") as archivo:
        assert {f["id_silabo"] for f in csv.DictReader(archivo)} == {"SIL_1", "SIL_MARKETING"}
    # El alcance TODAS no permite asignar a Marketing el texto oficial de Sistemas.
    aprobadas[1]["nombre_competencia"] = NOMBRE
    aprobadas[1]["descripcion_breve_competencia"] = DESCRIPCION
    with pytest.raises(ValueError, match="idénticos"):
        salida_catalogos.construir_catalogos_curriculares(
            registros,
            tmp_path / "rechazado",
            carrera="TODAS",
            periodo_academico="2026-2",
            competencias_tecnicas=aprobadas,
        )


def test_api_multicarrera_normaliza_aprueba_y_publica_un_solo_paquete(
    indice: Path, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    from agente.api import normalizador
    from agente.db.neo4j_catalogos import leer_catalogos
    from agente.normalizador.ejecuciones import GestorEjecuciones
    from api import servidor

    monkeypatch.setenv("NORMALIZADOR_CURRICULAR_LLM", "true")
    llamadas = _modelo(monkeypatch, [{"competencias": [_propuesta()]}])
    documento = Document()
    metadata = documento.add_table(rows=2, cols=2)
    metadata.cell(0, 0).text, metadata.cell(0, 1).text = "Curso", NOMBRE
    metadata.cell(1, 0).text, metadata.cell(1, 1).text = "Código", "10001"
    logro = documento.add_table(rows=2, cols=3)
    for i, texto in enumerate(("Logro de aprendizaje general", "Descripción", "Competencias")):
        logro.cell(0, i).text = texto
    logro.cell(1, 0).text, logro.cell(1, 1).text = "L1", LOGRO
    docx = BytesIO()
    documento.save(docx)
    fuente = BytesIO()
    with ZipFile(fuente, "w") as paquete:
        for carrera in ("INGENIERIA_DE_SISTEMAS", "MARKETING"):
            paquete.writestr(f"{carrera}/2026-2/Ciclo_01/curso.docx", docx.getvalue())
    gestor = GestorEjecuciones(tmp_path / "ejecuciones")
    futuros = []
    submit = gestor._executor.submit

    def enviar(*args, **kwargs):
        futuro = submit(*args, **kwargs)
        futuros.append(futuro)
        return futuro

    monkeypatch.setattr(gestor._executor, "submit", enviar)
    monkeypatch.setattr(normalizador, "gestor_ejecuciones", gestor)
    cliente = TestClient(servidor.app, client=("127.0.0.1", 0))
    respuesta = cliente.post(
        "/normalizador/silabos",
        files={"archivo": ("todas.zip", fuente.getvalue(), "application/zip")},
        data={"carrera": "TODAS", "periodo": "2026-2", "hitl": "0"},
    )
    assert respuesta.status_code == 202
    id_ejecucion = respuesta.json()["id_ejecucion"]
    assert len(futuros) == 1
    futuros[0].result(timeout=30)
    estado = cliente.get(f"/normalizador/ejecuciones/{id_ejecucion}").json()
    assert estado["estado"] in {"limpiado", "limpiado_con_advertencias"}, estado
    assert estado["limpieza_silabos"]["registros"] == 2
    assert estado["release_gate"]["decision"] == "ALLOW_IMPORT"
    assert len(llamadas) == 2
    objeto = gestor._obtener_objeto(id_ejecucion)
    with (objeto.directorio / "salidas/catalogo_habilidades.csv").open(
        encoding="utf-8-sig"
    ) as archivo:
        habilidades = list(csv.DictReader(archivo))
    assert len(habilidades) == 2
    assert {h["id_habilidad"] for h in habilidades} == {"HAB_TEC_001"}
    assert {h["nombre_habilidad"] for h in habilidades} == {NOMBRE, "Analítica comercial"}
    paquete_validado = leer_catalogos(objeto.directorio / "salidas")
    assert len(paquete_validado) == 6
    assert len(paquete_validado["catalogo_habilidades.csv"]) == 2
    staging = [
        json.loads(linea)
        for linea in (objeto.directorio / "limpios/silabos.jsonl").read_text().splitlines()
    ]
    assert len({r["id_curso"] for r in staging}) == 2
    assert {s["id_silabo"] for s in estado["progreso_llm"]["silabos"]} == {
        r["id_silabo"] for r in staging
    }
    assert len(list((tmp_path / "ejecuciones").glob("NOR_*"))) == 1


def test_respuesta_vacia_reenvia_silabo_hasta_habilidad_valida(
    indice: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    llamadas = _modelo(monkeypatch, [
        None,
        {"competencias": []},
        {"competencias": [_propuesta()]},
    ])
    auditoria: list[dict[str, object]] = []
    assert len(_inferir([_registro()], auditoria=auditoria)) == 1
    assert len(llamadas) == 3
    assert auditoria == []


def test_tercer_intento_puede_completar_habilidad(indice: Path, monkeypatch) -> None:
    llamadas = _modelo(
        monkeypatch,
        [{"competencias": []}] * 2 + [{"competencias": [_propuesta()]}],
    )
    assert len(_inferir([_registro()])) == 1
    assert len(llamadas) == 3


def test_tres_respuestas_vacias_omiten_silabo_y_continuan(indice: Path, monkeypatch) -> None:
    llamadas = _modelo(
        monkeypatch,
        [{"competencias": []}] * 3 + [{"competencias": [_propuesta()]}],
    )
    trazas = []
    auditoria = []
    propuestas = _inferir(
        [_registro(), _registro("SIL_2")],
        auditoria=auditoria,
        al_actualizar_progreso_silabo=trazas.append,
    )
    assert len(llamadas) == 4
    assert [propuesta["id_silabo"] for propuesta in propuestas] == ["SIL_2"]
    assert any(
        traza.id_silabo == "SIL_1"
        and traza.estado_analisis == "sin_propuesta"
        and traza.error_codigo == "HABILIDAD_MINIMA_NO_CUMPLIDA"
        for traza in trazas
    )
    assert auditoria == [{
        "codigo": "HABILIDAD_MINIMA_NO_CUMPLIDA",
        "severidad": "warning",
        "id_silabo": "SIL_1",
        "mensaje": "El sílabo no produjo una habilidad válida del catálogo tras 3 intentos.",
    }]


@pytest.mark.parametrize("habilitado", [False, True])
def test_cactus_docx_pipeline_y_cli_verifican_llm_y_habilidades(
    indice: Path, monkeypatch: pytest.MonkeyPatch, tmp_path: Path, habilitado: bool
) -> None:
    from agente.normalizador import ejecuciones
    from agente.normalizador.ejecuciones import GestorEjecuciones
    from agente.normalizador.silabos.fuente_cactus import ResultadoExtraccionCactus
    from scripts import normalizar_cactus_secuencial as cli

    monkeypatch.setenv("NORMALIZADOR_CURRICULAR_LLM", str(habilitado).lower())
    llamadas = _modelo(monkeypatch, [{"competencias": [_propuesta()]}])

    class Extractor:
        def __init__(self, **_kwargs):
            pass

        def extraer(self, **kwargs):
            carpeta = kwargs["directorio_salida"] / "Ciclo_01"
            carpeta.mkdir(parents=True)
            archivo = carpeta / "curso.docx"
            documento = Document()
            metadata = documento.add_table(rows=2, cols=2)
            metadata.cell(0, 0).text, metadata.cell(0, 1).text = "Curso", NOMBRE
            metadata.cell(1, 0).text, metadata.cell(1, 1).text = "Código", "10001"
            logro = documento.add_table(rows=2, cols=3)
            for i, texto in enumerate(
                ("Logro de aprendizaje general", "Descripción", "Competencias")
            ):
                logro.cell(0, i).text = texto
            logro.cell(1, 0).text, logro.cell(1, 1).text = "L1", LOGRO
            documento.save(archivo)
            return ResultadoExtraccionCactus(
                carrera=CARRERA,
                periodo="2026-2",
                cursos_encontrados=1,
                archivos_descargados=1,
                archivos_procesables=1,
                sin_silabo=0,
                fetch_fallidos=0,
                sesiones_fallidas=0,
                archivos_no_soportados=0,
                archivos=(archivo,),
                errores=(),
            )

    gestor = GestorEjecuciones(tmp_path / "ejecuciones")
    identificador, directorio = gestor.crear(
        "silabos",
        "cactus.zip",
        {"carrera": CARRERA, "periodo": "2026-2", "fuente": "cactus", "hitl": "0"},
    )
    monkeypatch.setattr(ejecuciones, "CactusExtractor", Extractor)
    gestor._extraer_y_validar_silabos(
        gestor._obtener_objeto(identificador), CARRERA, "2026-2", "usuario", "contraseña"
    )
    estado = gestor.obtener(identificador)
    with (directorio / "salidas/catalogo_habilidades.csv").open(encoding="utf-8-sig") as fuente:
        habilidades = list(csv.DictReader(fuente))
    monkeypatch.setattr(cli, "CARRERAS_ULIMA", (CARRERA,))
    monkeypatch.setattr(cli, "solicitar", lambda *_args, **_kwargs: estado)
    if habilitado:
        assert estado["release_gate"]["decision"] == "ALLOW_IMPORT", estado
        assert len(llamadas) == 1
        assert len(habilidades) == 1 and habilidades[0]["nombre_habilidad"] == NOMBRE
        assert cli.ejecutar("http://localhost", "2026-2", "u", "p", tmp_path / "cli") == 0
    else:
        assert llamadas == [] and habilidades == []
        with pytest.raises(RuntimeError, match="LLM desactivado"):
            cli.ejecutar("http://localhost", "2026-2", "u", "p", tmp_path / "cli")


def test_sin_candidatos_o_logros_analiza_sin_inventar_habilidades(
    indice: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    llamadas = _modelo(monkeypatch, [{"competencias": []}])
    registro = _registro()
    registro["carrera"] = "Carrera sin catálogo"
    assert _inferir([registro], exigir_habilidad=False) == []
    registro = _registro()
    registro["datos"]["logro_general"] = ""
    assert _inferir([registro], exigir_habilidad=False) == []
    assert len(llamadas) == 2


def test_comprobacion_usa_http_real_del_analista_y_embedding(indice, monkeypatch):
    import threading
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

    from api import servidor

    peticiones = []

    class Ollama(BaseHTTPRequestHandler):
        def log_message(self, *_args):
            pass

        def do_POST(self):
            datos = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            peticiones.append((self.path, datos["model"]))
            if self.path == "/api/embed":
                assert datos["model"] == "modelo-del-indice"
                respuesta = {"embeddings": [[1, 0]]}
            else:
                assert self.path == "/v1/chat/completions"
                assert datos["model"] == "modelo-analista-prueba"
                assert datos["response_format"]["type"] == "json_schema"
                respuesta = {
                    "id": "prueba",
                    "model": datos["model"],
                    "created": 0,
                    "object": "chat.completion",
                    "choices": [
                        {
                            "index": 0,
                            "finish_reason": "stop",
                            "message": {
                                "role": "assistant",
                                "content": '{"competencias": []}',
                            },
                        }
                    ],
                }
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(json.dumps(respuesta).encode())

    proveedor = ThreadingHTTPServer(("127.0.0.1", 0), Ollama)
    hilo = threading.Thread(target=proveedor.serve_forever, daemon=True)
    hilo.start()
    base = f"http://127.0.0.1:{proveedor.server_port}"
    monkeypatch.setenv("NORMALIZADOR_CURRICULAR_LLM", "true")
    monkeypatch.setenv("NORMALIZADOR_CURRICULAR_LLM_PROVIDER", "ollama")
    monkeypatch.setenv("NORMALIZADOR_CURRICULAR_OLLAMA_MODEL", "modelo-analista-prueba")
    monkeypatch.setenv("NORMALIZADOR_CURRICULAR_OLLAMA_BASE_URL", base + "/v1")
    monkeypatch.setenv("NORMALIZADOR_CURRICULAR_LLM_MAX_RETRIES", "0")
    monkeypatch.setenv("NORMALIZADOR_HAB_TEC_EMBEDDING_ENDPOINT", base + "/api/embed")
    monkeypatch.setattr(HabTecRetriever, "_embedding", EMBEDDING_REAL)
    try:
        respuesta = TestClient(servidor.app, client=("127.0.0.1", 0)).post(
            cli_ruta_comprobacion(), json={"carreras": [CARRERA, "Marketing"]}
        )
    finally:
        proveedor.shutdown()
        proveedor.server_close()
        hilo.join()
    assert respuesta.status_code == 200, respuesta.text
    assert respuesta.json()["catalogo"]["id_catalogo"] == indice.name
    assert respuesta.json()["catalogo"]["carreras"] == 2
    assert respuesta.json()["prueba_llm"]["latencia_modelo_ms"] is not None
    assert peticiones == [
        ("/api/embed", "modelo-del-indice"),
        ("/v1/chat/completions", "modelo-analista-prueba"),
    ]


def cli_ruta_comprobacion():
    from scripts.normalizar_cactus_secuencial import RUTA_COMPROBACION

    return RUTA_COMPROBACION


@pytest.mark.parametrize("carreras", [[], [CARRERA, CARRERA], ["Carrera inexistente"]])
def test_comprobacion_rechaza_lista_de_carreras_invalida(carreras):
    from api import servidor

    respuesta = TestClient(servidor.app, client=("127.0.0.1", 0)).post(
        cli_ruta_comprobacion(), json={"carreras": carreras}
    )
    assert respuesta.status_code == 422


@pytest.mark.parametrize("fallo", ["catalogo", "embedding", "analista"])
def test_comprobacion_falla_por_etapa_sin_filtrar_secretos(indice, monkeypatch, fallo):
    from agente.api import normalizador
    from api import servidor

    monkeypatch.setenv("NORMALIZADOR_CURRICULAR_LLM", "true")
    if fallo != "catalogo":
        monkeypatch.setattr(normalizador, "CARRERAS_ULIMA", (CARRERA, "Marketing"))

    def fallar(*_args, **_kwargs):
        raise RuntimeError("secreto-no-mostrar")

    monkeypatch.setattr(
        HabTecRetriever, "_embedding", fallar if fallo == "embedding" else lambda *_args: (1, 0)
    )
    monkeypatch.setattr(analista_tecnico, "obtener_llm", fallar)
    respuesta = TestClient(servidor.app, client=("127.0.0.1", 0)).post(cli_ruta_comprobacion())
    assert respuesta.status_code == 503
    assert (
        respuesta.json()["detail"]["codigo"]
        == {
            "catalogo": "CATALOGO_NO_DISPONIBLE",
            "embedding": "EMBEDDING_NO_DISPONIBLE",
            "analista": "LLM_NO_DISPONIBLE",
        }[fallo]
    )
    assert "secreto-no-mostrar" not in respuesta.text


def test_sin_indice_falla_sin_inventar_catalogo(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("NORMALIZADOR_HAB_TEC_DIR", str(tmp_path))
    with pytest.raises(ErrorHabTecRetriever, match="No hay un catálogo"):
        _inferir([_registro()])


def test_evidencia_inexistente_es_rechazada(indice: Path, monkeypatch) -> None:
    propuesta = _propuesta()
    propuesta["evidencia"][0]["fragmento"] = "Una evidencia ajena al sílabo."
    _modelo(monkeypatch, [{"competencias": [propuesta]}])
    assert _inferir([_registro()], exigir_habilidad=False) == []


def test_cancelacion_detiene_reintentos_sin_habilidad(indice: Path, monkeypatch) -> None:
    llamadas = _modelo(monkeypatch, [{"competencias": []}])
    with pytest.raises(CancelacionSolicitada):
        _inferir([_registro()], cancelada=lambda: len(llamadas) >= 2)
    assert len(llamadas) == 2


def test_cancelacion_se_respeta_entre_silabos(indice: Path, monkeypatch) -> None:
    llamadas = _modelo(monkeypatch, [{"competencias": [_propuesta()]}])
    cancelada = False

    def progreso(traza: Any) -> None:
        nonlocal cancelada
        if traza.estado_analisis == "completado":
            cancelada = True

    with pytest.raises(CancelacionSolicitada):
        _inferir(
            [_registro(), _registro("SIL_2")],
            cancelada=lambda: cancelada,
            al_actualizar_progreso_silabo=progreso,
        )
    assert len(llamadas) == 1


@pytest.mark.parametrize(
    "campo,valor",
    [
        ("catalogo_ref", "HAB_TEC_999"),
        ("catalogo_ref", ""),
        ("nombre_competencia", "arquitectura de software"),
        ("descripcion_breve_competencia", "Descripción editada."),
        ("catalogo_sha256", "sha-alterado"),
    ],
)
def test_exportacion_rechaza_campos_alterados_aun_aprobados(
    indice: Path, monkeypatch, tmp_path: Path, campo: str, valor: str
) -> None:
    _modelo(monkeypatch, [{"competencias": [_propuesta()]}])
    propuesta = _inferir([_registro()])[0]
    propuesta[campo] = valor
    salida = tmp_path / "salida"
    with pytest.raises(ValueError):
        salida_catalogos.construir_catalogos_curriculares(
            [_registro()],
            salida,
            carrera="INGENIERIA_DE_SISTEMAS",
            periodo_academico="2026-2",
            competencias_tecnicas=[{**propuesta, "estado_aprobacion": "APROBADA"}],
        )
    assert not (salida / "catalogo_habilidades.csv").exists()


def test_exportacion_rechaza_habilidad_de_otra_carrera(indice: Path, monkeypatch) -> None:
    catalogo = analista_tecnico.cargar_catalogo_hab_tec()
    propuesta = {
        **_propuesta(),
        "nombre_competencia": "Analítica comercial",
        "descripcion_breve_competencia": "Analiza campañas.",
    }
    with pytest.raises(ValueError, match="idénticos"):
        analista_tecnico.validar_habilidad_catalogo(propuesta, CARRERA, catalogo)


def test_catalogo_nuevo_no_cambia_version_de_propuesta_aprobada(
    indice: Path, monkeypatch, tmp_path: Path
) -> None:
    _modelo(monkeypatch, [{"competencias": [_propuesta()]}])
    propuesta = _inferir([_registro()])[0]
    nuevo = _vectorizar(indice.parent, monkeypatch)
    assert nuevo != indice
    assert cargar_retriever_hab_tec().id_catalogo == nuevo.name
    candidato = analista_tecnico.validar_habilidad_catalogo(propuesta, CARRERA)
    assert candidato.nombre == NOMBRE


def test_catalogo_configurado_prevalece_sobre_el_mas_reciente(
    indice: Path, monkeypatch
) -> None:
    nuevo = _vectorizar(indice.parent, monkeypatch)
    assert cargar_retriever_hab_tec().id_catalogo == nuevo.name
    monkeypatch.setenv("NORMALIZADOR_HAB_TEC_ID", indice.name)
    assert cargar_retriever_hab_tec().id_catalogo == indice.name
    assert cargar_retriever_hab_tec(nuevo.name).id_catalogo == nuevo.name


@pytest.mark.parametrize(
    "archivo", ["catalogo_normalizado.jsonl", "indice/metadatos.jsonl", "indice/vectores.f32"]
)
def test_indice_inconsistente_bloquea_inferencia(indice: Path, archivo: str) -> None:
    ruta = indice / archivo
    ruta.write_bytes(ruta.read_bytes() + b"alterado")
    with pytest.raises(ErrorHabTecRetriever):
        analista_tecnico.cargar_catalogo_hab_tec()


def test_catalogo_con_vectorizacion_fallida_no_se_activa(indice: Path) -> None:
    ruta = indice / "manifest.json"
    manifest = json.loads(ruta.read_text(encoding="utf-8"))
    manifest["estado"] = "error"
    ruta.write_text(json.dumps(manifest), encoding="utf-8")
    assert cargar_retriever_hab_tec() is None
    with pytest.raises(ErrorHabTecRetriever):
        cargar_retriever_hab_tec(indice.name)


def test_hash_literal_del_catalogo_se_conserva_hasta_importacion(
    indice: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from agente.db.neo4j_catalogos import leer_catalogos

    id_habilidad = "f9db583c7660"
    _vectorizar(indice.parent, monkeypatch, id_habilidad)
    _modelo(monkeypatch, [{"competencias": [_propuesta(id_habilidad)]}])
    registro = _registro("1234567890abcdef")
    registro["id_curso"] = "CUR_1234567890abcdef"
    registro["id_silabo"] = "SIL_1234567890abcdef"
    propuestas = _inferir([registro])
    assert propuestas[0]["catalogo_ref"] == id_habilidad
    propuestas[0]["estado_aprobacion"] = "APROBADA"
    salida = tmp_path / "salida_hash"
    salida_catalogos.construir_catalogos_curriculares(
        [registro],
        salida,
        carrera="INGENIERIA_DE_SISTEMAS",
        periodo_academico="2026-2",
        competencias_tecnicas=propuestas,
    )
    filas = leer_catalogos(salida)
    assert filas["catalogo_habilidades.csv"][0]["id_habilidad"] == id_habilidad
    assert filas["catalogo_habilidades.csv"][0]["nombre_habilidad"] == NOMBRE
    assert filas["catalogo_habilidades.csv"][0]["desc_breve"] == DESCRIPCION
    assert {f["id_habilidad"] for f in filas["cobertura_curricular.csv"]} == {id_habilidad}
